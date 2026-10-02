"""
compare_tool.py - Check, clean, compare & pivot ANY tabular data (settings version)
==================================================================================

Works with any kind of data: sales, inventory, HR, finance, logistics, surveys,
customer lists, exported reports ... The tool looks at the values themselves,
so it does not need to know in advance what your columns mean.

HOW TO USE
  1. Edit the SETTINGS section below. Look for the 👉 comments.
  2. Run:      python compare_tool.py
  3. Open the Excel report saved at OUTPUT_FILE.

First time? Leave USE_DEMO_DATA = True and run it once per TASK to see what each one does.
When you switch to your own files, change the column names in the settings to yours -
if a name is wrong, the tool tells you which setting and suggests the right name.

Recommended order with new data:   check_data  ->  clean_data  ->  compare / pivot

Install once:  pip install pandas openpyxl rapidfuzz
               (+ xlrd for old .xls files, + pyarrow for .parquet files)

This file is also the engine for data_tool.py (the click version) - keep them in one folder.
"""

# ╔══════════════════════════════════════════════════════════════════════════╗
# ║                  ✏️  SETTINGS - EDIT ONLY THIS PART                       ║
# ╚══════════════════════════════════════════════════════════════════════════╝

# 👉 STEP 0 - Test with built-in example data first?
#    True  = ignore your files and use example data (good for a first try)
#    False = use YOUR files from STEP 2
USE_DEMO_DATA = True


# 👉 STEP 1 - What do you want to do? (pick ONE and write it between the quotes)
#    "check_data"      -> data quality report: empty cells, numbers saved as text, mixed date
#                         formats, spelling variants, possible typos, duplicates, odd values ...
#                         (nothing is changed - it only reports)
#    "clean_data"      -> clean the data using STEP 10 and save a cleaned copy + a log of every change
#    "fuzzy_compare"   -> match rows of TABLE_A and TABLE_B even when names are spelled
#                         differently, e.g. "Apex Cement Co." vs "APEX CEMENT FACTORY"
#    "compare_tables"  -> match rows by a key that is EXACTLY the same in both tables
#                         (invoice no., item code, employee ID ...)
#    "compare_columns" -> compare one column with another column
#    "compare_rows"    -> compare one row with another row
#    "pivot"           -> build a pivot table from TABLE_A
TASK = "check_data"


# 👉 STEP 2 - Your files: CSV, TXT, TSV, XLSX, XLSM, XLS, JSON or Parquet.
#    Power BI: on a visual click "..." -> Export data, then put that file here.
#    Windows paths: use forward slashes    "C:/Users/Me/Desktop/sales.xlsx"
#                   or put r in front      r"C:\Users\Me\Desktop\sales.xlsx"
#    "sheet" = the Excel sheet (tab) name. None = the first sheet. Ignored for CSV.
#    To compare two sheets of the SAME Excel file, use the same "file" with a different "sheet".
#    TABLE_B is not used by "pivot". For "check_data"/"clean_data" it is optional
#    (see INCLUDE_TABLE_B in STEP 10, or set "file": "" to skip it).
TABLE_A = {"file": "data/erp_suppliers.xlsx",      "sheet": None}
TABLE_B = {"file": "data/supplier_statement.csv",  "sheet": None}


# 👉 STEP 3 - KEY column(s): the column that says WHICH row is which
#    (customer name, supplier name, invoice no., item code, employee ID ...).
#    Write the names EXACTLY as they appear in the header row of each file.
#    The names may be different in the two tables.
#    More than one column is allowed. Keep the same order on both sides:
#        KEY_A = ["Customer", "City"]
#        KEY_B = ["Client",   "Town"]
#    Used by: "fuzzy_compare" and "compare_tables"
KEY_A = ["Supplier Name"]
KEY_B = ["Vendor"]


# 👉 STEP 4 - Which columns should be checked once rows are matched?
#    Format:  {"column name in TABLE_A": "column name in TABLE_B", ...}
#    Leave it empty  {}  to check every column that has the same name in both tables.
#    Used by: "fuzzy_compare", "compare_tables" and "compare_rows"
COMPARE_COLUMNS = {"Amount": "Total", "City": "City"}


# 👉 STEP 5 - Fuzzy settings (only for "fuzzy_compare")
SIMILARITY_THRESHOLD = 85     # 0-100. Pairs scoring below this are NOT matched.
                              # 85-90 is a good start. Lower = more matches, but more wrong ones.
REVIEW_BELOW = 95             # Matches scoring below this get "Needs review" (yellow in Excel).

MATCH_METHOD = "token_sort"   # "token_sort" -> ignores word order. Best for names.
                              #                 "Steel Blue River" = "Blue River Steel"
                              # "token_set"  -> OK when one side has extra words.
                              #                 "Apex Cement" vs "Apex Cement Factory Baghdad"
                              # "simple"     -> letter-by-letter. Best for codes/IDs with typos.
                              # "partial"    -> one text is inside the other.

IGNORE_WORDS = ["co", "company", "ltd", "llc", "inc", "the"]
                              # Words removed before comparing names (case doesn't matter). [] = none.

MUST_MATCH_EXACTLY = {}       # Optional safety rule: only pair rows that ALSO have exactly
                              # the same value in these columns.
                              # Example: {"Invoice Date": "Date"}   or   {"City": "City"}
                              # {} = off


# 👉 STEP 6 - General comparison options
NUMBER_TOLERANCE = 0.0        # Numbers count as "same" if they differ by no more than this.
                              # Example: 0.01 to ignore rounding differences.
IGNORE_CASE = True            # True = "ERBIL" and "erbil" count as the same text.


# 👉 STEP 7 - Pivot settings (only for "pivot", uses TABLE_A)
PIVOT_ROWS = ["City"]         # fields down the left side
PIVOT_COLUMNS = ["Category"]  # fields across the top          ([] = none)
PIVOT_VALUES = ["Amount"]     # fields to calculate            ([] = just count rows)
PIVOT_CALC = "sum"            # "sum", "mean", "count", "nunique", "min", "max" or "median"
PIVOT_TOTALS = True           # add Total row/column


# 👉 STEP 8 - Column comparison (only for "compare_columns")
COLUMN_A = "Supplier Name"    # a column in TABLE_A
COLUMN_B = "Vendor"           # a column in TABLE_B
                              # (to compare two columns of ONE table, point TABLE_B to the same file)
COMPARE_BY_POSITION = False   # True = also compare row 1 with row 1, row 2 with row 2 ...
                              # Only useful if both tables are sorted the same way.


# 👉 STEP 9 - Row comparison (only for "compare_rows")
#    "find_by" = "row_number" (first data row is 0) or the name of a column to search in
#    "value"   = the row number, or the value to look for in that column
#    Columns with the same name are compared, plus the pairs you set in COMPARE_COLUMNS (STEP 4).
ROW_A = {"find_by": "Supplier Name", "value": "Apex Cement Factory"}
ROW_B = {"find_by": "Vendor",        "value": "Apex Cement Fctory"}


# 👉 STEP 10 - CLEANING (works on any kind of data)
#    Used by "clean_data". If APPLY_CLEANING = True it ALSO runs on both tables
#    before every other task, so comparisons and pivots use clean values.
#    Every change is listed in the "Cleaning log" sheet. Nothing changes silently,
#    and no rows are deleted except 100% identical duplicates (if you allow it below).
APPLY_CLEANING = True

EMPTY_MARKERS = ["", "-", "--", "n/a", "na", "null", "none", "nil", "#n/a", "?", "nan"]
                              # Cell values that really mean "empty". Upper/lower case doesn't matter.

TRIM_SPACES = True            # Remove spaces at the start/end, double spaces, and invisible characters
                              # (often copied from websites, PDFs or other systems).

UNIFY_SPELLING_VARIANTS = True
                              # "Riverside" / "RIVERSIDE" / "riverside " -> the most common spelling.
                              # Only fixes upper/lower case and spaces. Real typos ("Hilview") are
                              # listed by check_data in "Similar spellings" - fix those with REPLACE_VALUES.

TO_NUMBER = "auto"            # Turn text into real numbers so they can be summed and compared:
                              #   "auto"              = columns where almost all values look like numbers
                              #   ["Amount", "Price"] = only these columns        [] = never
                              # Understands: "1,250"  "$1,250.50"  "IQD 15,000"  "١٢٣٤" (Arabic digits)
                              #              "(500)" = -500   "500-" = -500   "15%" = 15
                              # Codes with leading zeros (phones, "00123") and very long IDs stay text.
DECIMAL_SEPARATOR = "."       # "." if your numbers look like 1,234.56   |  "," if they look like 1.234,56

TO_DATE = "auto"              # Turn text into real dates: "auto", ["Invoice Date"], or []
                              # Mixed formats are OK: "2026-03-01", "01/03/2026", "1 Mar 2026",
                              # and Excel date numbers like 46082.
DAY_FIRST = True              # True:  01/03/2026 = 1 March    (Iraq, Europe, most countries)
                              # False: 01/03/2026 = January 3  (USA)

FIX_CASE = {}                 # Force upper/lower/title case:  {"City": "title", "Code": "upper"}

REPLACE_VALUES = {}           # Fix known variants/typos. Matching ignores case and extra spaces.
                              # Example:
                              # {"City":   {"hawler": "Erbil", "slemani": "Sulaymaniyah"},
                              #  "Status": {"done": "Completed", "finished": "Completed"}}

REMOVE_DUPLICATE_ROWS = True  # Remove rows that are 100% identical (checked AFTER the fixes above).

MERGE_DUPLICATE_KEYS = {}     # Combine rows that share a key into one row.   {} = off
                              # Example: {"key": ["Invoice No"], "how": {"Amount": "sum", "Date": "max"}}
                              # "how" options: sum, mean, min, max, first, last, count,
                              #                join (puts the different texts together: "A, B")
                              # Columns not listed in "how" keep their first non-empty value.

INCLUDE_TABLE_B = True        # "check_data" / "clean_data" also process TABLE_B (False = only TABLE_A)
CLEANED_FOLDER = "results/cleaned"   # where "clean_data" saves the cleaned copies of your files


# 👉 STEP 11 - RULES your data should follow (optional, any kind of data).
#    Rows that break a rule are LISTED in the "Data problems" sheet - they are NOT deleted.
#    Rules are checked when cleaning runs (clean_data, or any task with APPLY_CLEANING = True).
#    Each rule is {"column": "...", <check>: <value>}. Checks you can use:
#      "required": True            -> must not be empty
#      "unique": True              -> no repeated values (IDs, invoice numbers, emails ...)
#      "min": 0   /  "max": 100    -> number range (dates too: "min": "2020-01-01")
#      "allowed": ["Open", "Closed"]  -> only these values (case doesn't matter)
#      "pattern": r"^\d{11}$"      -> must match a pattern. Examples:
#                                       r"^\d{11}$"                  exactly 11 digits
#                                       r"^[^@\s]+@[^@\s]+\.\w+$"    looks like an email
#      "not_future": True          -> date must not be after today
#      "not_after": "End Date"     -> value must not be bigger/later than another column
#    A rule for a column that exists in only one table is applied to that table only.
#    RULES = []  turns rules off.
RULES = [
    {"column": "Amount", "min": 0},
    {"column": "Invoice Date", "not_future": True},
    {"column": "Supplier Name", "required": True},
]


# 👉 STEP 12 - Where to save the report (the folder is created if needed)
OUTPUT_FILE = "results/result.xlsx"


# ╔══════════════════════════════════════════════════════════════════════════╗
# ║                ⛔  NO NEED TO EDIT ANYTHING BELOW THIS LINE               ║
# ╚══════════════════════════════════════════════════════════════════════════╝

import difflib
import io
import json
import os
import re
import sys
import warnings
from datetime import date, datetime

import numpy as np
import pandas as pd

try:
    from rapidfuzz import fuzz, process
except ImportError:  # only needed for fuzzy matching and the "similar spellings" check
    fuzz = process = None


FILE_TYPES = ["csv", "tsv", "txt", "xlsx", "xlsm", "xls", "json", "parquet"]
EXCEL_MAX_ROWS = 1_048_575

AGG_FUNCS = {"Sum": "sum", "Average": "mean", "Count": "count", "Distinct count": "nunique",
             "Min": "min", "Max": "max", "Median": "median"}
NUMERIC_AGGS = {"sum", "mean", "median"}

MATCH_METHODS = {
    "token_sort": ("token_sort_ratio", "Names - ignores word order"),
    "token_set": ("token_set_ratio", "Names with extra words on one side"),
    "simple": ("ratio", "Codes / IDs with typos"),
    "partial": ("partial_ratio", "One text contained in the other"),
}

RULE_CHECKS = {"required", "unique", "min", "max", "allowed", "pattern", "not_future", "not_after"}
MERGE_HOW = {"sum", "mean", "min", "max", "first", "last", "count", "join"}
CASE_MODES = {"upper", "lower", "title"}

TYPE_LABELS = {"empty": "empty", "number": "number", "number_text": "number (stored as text)",
               "date": "date", "date_text": "date (stored as text)", "code": "code / ID (text)",
               "yesno": "yes/no", "yesno_text": "yes/no (text)", "text": "text"}

# Arabic-Indic and Persian/Kurdish digits + Arabic decimal/thousands separators -> western
ARABIC_DIGITS = str.maketrans("٠١٢٣٤٥٦٧٨٩۰۱۲۳۴۵۶۷۸۹٫٬", "01234567890123456789.,")
INVISIBLE = "[\u200b\u200e\u200f\u202a-\u202e\u2066-\u2069\ufeff]"   # (keeps \u200c, used in Kurdish/Persian)
CURRENCY = (r"(?i)\b(?:iqd|usd|eur|gbp|try|aed|sar|irr|dinars?|dollars?|euros?)\b"
            r"|د\.?\s?ع\.?|[$€£¥₺﷼]")
NUMBER_SHAPE = r"^[+-]?\(?[+-]?\d[\d.,]*\)?-?%?$"
DATE_LIKE = (r"(?i)^\s*(?:\d{1,4}[-/.]\d{1,2}[-/.]\d{1,4}"
             r"|\d{1,2}[\s-]+(?:jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*\.?[\s,-]+\d{2,4}"
             r"|(?:jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*\.?\s+\d{1,2},?\s+\d{2,4})")
YESNO = {"yes", "no", "true", "false", "y", "n", "نعم", "لا", "بەڵێ", "نەخێر"}
ID_HINT = r"(?i)(^id$|_id$|\bid\b|\bno\.?$|\bnumber\b|\bcode\b|رقم|ژمارە)"

DEFAULT_CLEANING = {
    "empty_markers": ["", "-", "--", "n/a", "na", "null", "none", "nil", "#n/a", "?", "nan"],
    "trim": True, "unify": True, "to_number": "auto", "decimal": ".", "to_date": "auto",
    "day_first": True, "fix_case": {}, "replace": {}, "remove_dups": True, "merge": {}, "rules": [],
}


class SettingsError(Exception):
    """A problem the user can fix in the SETTINGS section."""


# =============================================================================
# 1. LOADING
# =============================================================================
PLAIN_NUMBER = r"^[+-]?(\d+\.?\d*|\.\d+)([eE][+-]?\d+)?$"


def _infer_csv_types(df: pd.DataFrame) -> pd.DataFrame:
    """Turn plain-number text columns into numbers, but keep codes with leading zeros as text."""
    for c in df.columns:
        v = df[c].dropna().pipe(as_text).str.strip()
        if v.empty:
            continue
        if v.str.match(PLAIN_NUMBER).all() and not v.str.match(r"^[+-]?0\d").any():
            num = pd.to_numeric(df[c].pipe(as_text).str.strip().where(df[c].notna()), errors="coerce")
            whole = num.dropna()
            if len(whole) and (whole % 1 == 0).all() and whole.abs().max() < 2**53:
                num = num.astype("Int64")
            df[c] = num
    return df


def _read_text(data: bytes, sep):
    """Read CSV/TXT/TSV as text first (so '0750...' keeps its zero), then detect number columns.
    Tries common encodings, including Arabic Windows (cp1256)."""
    for enc in ("utf-8-sig", "cp1256", "latin-1"):
        try:
            df = pd.read_csv(io.BytesIO(data), sep=sep, engine="python", encoding=enc, dtype=str)
            return _infer_csv_types(df)
        except UnicodeDecodeError:
            continue
    raise ValueError("Could not detect the text encoding of this file.")


def clean_df(df: pd.DataFrame) -> pd.DataFrame:
    """Basic tidy on load: text column names, drop fully empty rows/columns."""
    df = df.copy()
    df.columns = [str(c).strip() for c in df.columns]
    df = df.dropna(how="all").dropna(axis=1, how="all")
    return df.reset_index(drop=True)


def read_file_bytes(filename: str, data: bytes) -> dict:
    """Return {sheet_name: DataFrame}. Non-Excel files use the key ""."""
    ext = os.path.splitext(filename)[1].lower().lstrip(".")
    if ext in ("csv", "txt"):
        return {"": clean_df(_read_text(data, sep=None))}      # auto-detects , ; | tab
    if ext == "tsv":
        return {"": clean_df(_read_text(data, sep="\t"))}
    if ext in ("xlsx", "xlsm", "xls"):
        # dtype=object keeps cells exactly as typed in Excel (text stays text, e.g. '0750...');
        # infer_objects then gives real number/date cells their proper type.
        sheets = pd.read_excel(io.BytesIO(data), sheet_name=None, dtype=object)
        return {str(s): clean_df(df).infer_objects() for s, df in sheets.items()}
    if ext == "json":
        return {"": clean_df(pd.read_json(io.BytesIO(data)))}
    if ext == "parquet":
        return {"": clean_df(pd.read_parquet(io.BytesIO(data)))}
    raise ValueError(f"Unsupported file type '.{ext}'. Supported: {', '.join(FILE_TYPES)}")


def load_table(cfg: dict, setting: str) -> pd.DataFrame:
    """Load one table from a TABLE_A / TABLE_B setting."""
    path = str(cfg.get("file", "")).strip()
    if not path:
        raise SettingsError(f'{setting}["file"] is empty. Put the path to your file there.')
    if not os.path.isfile(path):
        raise SettingsError(f'File not found: "{path}"  (check {setting}["file"])\n'
                            f'  Current folder is: {os.getcwd()}')
    with open(path, "rb") as f:
        try:
            sheets = read_file_bytes(path, f.read())
        except ValueError as e:
            raise SettingsError(f"{setting}: {e}")
    sheet = cfg.get("sheet")
    if "" in sheets:
        return sheets[""]
    if sheet is None:
        return next(iter(sheets.values()))
    if str(sheet) not in sheets:
        raise SettingsError(f'Sheet "{sheet}" not found in "{path}" (check {setting}["sheet"]).\n'
                            f"  Sheets in this file: {list(sheets)}")
    return sheets[str(sheet)]


def demo_tables():
    """Example data with typical real-world problems (any field has the same kinds of issues)."""
    a = pd.DataFrame({
        "Supplier Name": ["Northstar Trading Co.", "Blue River Steel Ltd", "Apex Cement Factory",
                          "Global Logistics LLC", "Silver Line Electronics", "Green Valley Foods",
                          "Summit Office Supplies", "Delta Fuel Services", "Northstar Trading Co.  "],
        "City": ["Riverside", "Hillview", "Lakeside", "riverside ", "Hillview", "LAKESIDE",
                 "Riverside", "Hilview", "Riverside"],
        "Category": ["Trading", "Construction", "Construction", "Transport", "IT", "Food",
                     "Office", "Fuel", "Trading"],
        "Amount": ["IQD 12,500", "48,200.50", "٣١٠٠٠", "9800", "$15,250", "7,400", "(2,300)",
                   "18900", "IQD 12,500"],
        "Invoice Date": ["2026-03-01", "01/03/2026", "5 Mar 2026", "2026-03-07", "N/A",
                         "12/03/2026", "2030-01-15", "2026-03-20", "2026-03-01"],
        "Phone": ["07501234567", "07701112233", "-", "07509998877", "07501231234",
                  "07705556677", "07501112222", "07703334444", "07501234567"],
    })
    b = pd.DataFrame({
        "Vendor": ["NORTHSTAR TRADING COMPANY", "Blue River Steel", "Apex Cement Fctory",
                   "Global Logistics", "Silverline Electronics", "Green Valley Food Co",
                   "Orion Medical Supplies", "Delta Fuels Services"],
        "City": ["Riverside", "Hillview", "Lakeside", "Riverside", "Hillview", "Hillview",
                 "Lakeside", "Hillview"],
        "Total": [12500, 48200.5, 31500, 9800, 15250, 7400, 5600, 18900],
    })
    return a, b


# =============================================================================
# 2. SMALL HELPERS
# =============================================================================
def _suggest(name, available):
    guess = difflib.get_close_matches(str(name), [str(a) for a in available], n=1, cutoff=0.6)
    return f'  -> did you mean "{guess[0]}"?' if guess else ""


def check_columns(df: pd.DataFrame, cols, table: str, setting: str):
    """Friendly error (with 'did you mean') if a column name is wrong."""
    available = [str(c) for c in df.columns]
    problems = [f'  - "{c}" is not a column in {table}{_suggest(c, available)}'
                for c in cols if c not in df.columns]
    if problems:
        raise SettingsError(f"Problem in setting {setting}:\n" + "\n".join(problems)
                            + f"\n  Columns in {table}: {available}")


def as_text(s):
    """Text as plain Python strings, so the same regex rules work on every pandas version."""
    return s.astype(str).astype(object)


def _str_mask(s: pd.Series) -> pd.Series:
    return s.map(lambda x: isinstance(x, str)).astype(bool)


def _num_mask(s: pd.Series) -> pd.Series:
    return s.map(lambda x: isinstance(x, (int, float, np.integer, np.floating))
                 and not isinstance(x, (bool, np.bool_))).astype(bool)


def is_text_col(s: pd.Series) -> bool:
    return pd.api.types.is_object_dtype(s.dtype) or isinstance(s.dtype, pd.StringDtype)


def tidy_spaces(t: pd.Series) -> pd.Series:
    """Remove invisible characters, collapse spaces, trim."""
    return (t.str.replace(INVISIBLE, "", regex=True)
             .str.replace(r"\s+", " ", regex=True).str.strip())


def map_text(s: pd.Series, func) -> pd.Series:
    """Apply func (Series -> Series) to the text cells only; numbers/dates stay untouched."""
    m = _str_mask(s)
    if not m.any():
        return s
    out = s.astype(object).copy()
    out.loc[m] = func(s[m].pipe(as_text)).to_numpy()
    return out


def norm_text(s: pd.Series, ignore_case: bool = False) -> pd.Series:
    """Comparable text: trims spaces, '1001.0' -> '1001', dates -> 'YYYY-MM-DD', empty -> ''."""
    if pd.api.types.is_datetime64_any_dtype(s):
        out = s.dt.strftime("%Y-%m-%d %H:%M:%S").astype(object).str.replace(" 00:00:00", "", regex=False)
    else:
        out = s.pipe(as_text).str.strip().str.replace(r"\.0$", "", regex=True)
    out = out.where(s.notna(), "")
    return out.str.lower() if ignore_case else out


def values_equal(a, b, tol: float = 0.0, ignore_case: bool = False) -> np.ndarray:
    """Cell-by-cell equality. Numbers use a tolerance, text is trimmed."""
    a = pd.Series(a).reset_index(drop=True)
    b = pd.Series(b).reset_index(drop=True)
    both_empty = (a.isna() & b.isna()).astype(bool)
    one_empty = (a.isna() ^ b.isna()).astype(bool)
    dt_a, dt_b = pd.api.types.is_datetime64_any_dtype(a), pd.api.types.is_datetime64_any_dtype(b)
    if dt_a or dt_b:
        num_a = pd.to_datetime(a, errors="coerce") if dt_a else pd.Series(pd.NaT, index=a.index)
        num_b = pd.to_datetime(b, errors="coerce") if dt_b else pd.Series(pd.NaT, index=b.index)
        both_num = (num_a.notna() & num_b.notna()).astype(bool)
        num_eq = (num_a == num_b).fillna(False).astype(bool)
    else:
        num_a = pd.to_numeric(a, errors="coerce").astype(float)
        num_b = pd.to_numeric(b, errors="coerce").astype(float)
        both_num = (num_a.notna() & num_b.notna()).astype(bool)
        num_eq = ((num_a - num_b).abs() <= tol).fillna(False).astype(bool)
    txt_eq = (norm_text(a, ignore_case) == norm_text(b, ignore_case)).fillna(False).astype(bool)
    return (both_empty | (both_num & num_eq) | (~both_num & ~one_empty & txt_eq)).to_numpy()


def flatten_columns(df: pd.DataFrame) -> pd.DataFrame:
    if isinstance(df.columns, pd.MultiIndex):
        df.columns = [" | ".join(str(x) for x in col if str(x) != "") for col in df.columns]
    else:
        df.columns = [str(c) for c in df.columns]
    return df


def file_row(i):
    """Row number as seen in Excel (header is row 1, first data row is 2)."""
    return np.asarray(i) + 2


def _show(v):
    """Dates without a time part are shown as plain dates."""
    if isinstance(v, pd.Timestamp) and v == v.normalize():
        return v.date()
    return v


def _examples(values, n=3) -> str:
    vals = list(dict.fromkeys(str(v) for v in values))[:n]
    return " | ".join(vals)


# =============================================================================
# 3. UNDERSTANDING VALUES (numbers, dates, types)
# =============================================================================
def parse_numbers(values: pd.Series, decimal: str = "."):
    """Read numbers from any column. Returns (float Series, failed mask)."""
    s = values
    if pd.api.types.is_bool_dtype(s):
        return s.astype(float), pd.Series(False, index=s.index)
    if pd.api.types.is_numeric_dtype(s):
        return pd.to_numeric(s, errors="coerce").astype(float), pd.Series(False, index=s.index)
    out = pd.Series(np.nan, index=s.index, dtype=float)
    is_num = _num_mask(s)
    if is_num.any():
        out[is_num] = s[is_num].astype(float)
    strm = _str_mask(s)
    if strm.any():
        t = s[strm].pipe(as_text).str.translate(ARABIC_DIGITS)
        t = t.str.replace(CURRENCY, "", regex=True).str.replace(r"\s+", "", regex=True)
        shape_ok = t.str.match(NUMBER_SHAPE).fillna(False).astype(bool)
        core = t.str.replace("%", "", regex=False)
        neg = (core.str.match(r"^\+?\(|^-").fillna(False) | core.str.endswith("-").fillna(False)).astype(bool)
        digits = core.str.replace(r"[^\d.,]", "", regex=True)
        if decimal == ",":
            valid = digits.str.match(r"^(\d{1,3}(\.\d{3})+|\d+)(,\d+)?$")
            digits = digits.str.replace(".", "", regex=False).str.replace(",", ".", regex=False)
        else:
            valid = digits.str.match(r"^(\d{1,3}(,\d{3})+|\d+)(\.\d+)?$")
            digits = digits.str.replace(",", "", regex=False)
        ok = shape_ok & valid.fillna(False).astype(bool)
        val = pd.to_numeric(digits.where(ok), errors="coerce").astype(float)
        out[strm] = val.where(~neg, -val)
    failed = (s.notna() & out.isna()).astype(bool)
    return out, failed


def _one_date(x, day_first):
    try:
        if isinstance(x, str) and re.match(r"^\s*\d{4}[-/.]\d{1,2}[-/.]\d{1,2}", x):
            day_first = False
        ts = pd.to_datetime(x, dayfirst=day_first)
        return ts.tz_localize(None) if ts.tzinfo else ts
    except Exception:
        return pd.NaT


def parse_dates(values: pd.Series, day_first: bool = True):
    """Read dates from any column (mixed formats, Excel serial numbers). Returns (dates, failed)."""
    s = values
    if pd.api.types.is_datetime64_any_dtype(s):
        return pd.to_datetime(s), pd.Series(False, index=s.index)
    out = pd.Series(pd.NaT, index=s.index, dtype="datetime64[ns]")
    is_dt = s.map(lambda x: isinstance(x, (pd.Timestamp, datetime, date))).astype(bool)
    if is_dt.any():
        out[is_dt] = pd.to_datetime(s[is_dt].map(lambda x: _one_date(x, day_first)))
    is_num = _num_mask(s)
    if is_num.any():                                    # Excel stores dates as day numbers
        nums = s[is_num].astype(float)
        ok = nums.between(20000, 80000)
        if ok.any():
            out[nums[ok].index] = pd.to_datetime(nums[ok], unit="D", origin="1899-12-30")
    strm = _str_mask(s)
    if strm.any():
        t = s[strm].pipe(as_text).str.translate(ARABIC_DIGITS).str.strip()
        # Year-first dates (2026-03-01) are never day-first; parse them separately
        iso = t.str.match(r"^\d{4}[-/.]\d{1,2}[-/.]\d{1,2}").fillna(False).astype(bool)
        for part, dayfirst in ((t[iso], False), (t[~iso], day_first)):
            if part.empty:
                continue
            with warnings.catch_warnings():
                warnings.simplefilter("ignore")
                try:
                    parsed = pd.to_datetime(part, format="mixed", dayfirst=dayfirst, errors="coerce")
                    if getattr(parsed.dt, "tz", None) is not None:
                        parsed = parsed.dt.tz_localize(None)
                except Exception:
                    parsed = pd.to_datetime(part.map(lambda x: _one_date(x, dayfirst)))
            out[part.index] = parsed
    failed = (s.notna() & out.isna()).astype(bool)
    return out, failed


def detect_kind(values: pd.Series, day_first=True, decimal=".") -> str:
    """Work out what a column really contains (empty/number/date/code/yes-no/text)."""
    v = values.dropna()
    if v.empty:
        return "empty"
    if pd.api.types.is_bool_dtype(v):
        return "yesno"
    if pd.api.types.is_datetime64_any_dtype(v):
        return "date"
    if pd.api.types.is_numeric_dtype(v):
        return "number"
    strs = v[_str_mask(v)].pipe(as_text).str.strip()
    if strs.empty:
        if v.map(lambda x: isinstance(x, (pd.Timestamp, datetime, date))).all():
            return "date"
        return "number" if not parse_numbers(v, decimal)[1].any() else "text"
    t = strs.str.translate(ARABIC_DIGITS)
    if t.str.match(r"^0\d+$").mean() >= 0.3 or t.str.match(r"^\d{16,}$").mean() >= 0.3:
        return "code"
    if t.str.contains(DATE_LIKE, regex=True).mean() >= 0.8:
        if parse_dates(v, day_first)[1].mean() <= 0.1:
            return "date_text"
    if parse_numbers(v, decimal)[1].mean() <= 0.1:
        return "number_text"
    if t.str.lower().isin(YESNO).mean() >= 0.95:
        return "yesno_text"
    return "text"


def similar_values(texts: pd.Series, min_score: int = 88, limit: int = 30) -> list:
    """Pairs of different values that look like typos of each other ('Hilview' ~ 'Hillview')."""
    if process is None or texts.empty:
        return []
    vc = tidy_spaces(texts.pipe(as_text)).value_counts()
    d = pd.DataFrame({"v": vc.index, "n": vc.to_numpy(), "k": vc.index.str.lower()})
    agg = d.groupby("k", sort=False).agg(v=("v", "first"), n=("n", "sum"))
    agg = agg.sort_values("n", ascending=False).head(3000)
    keys = [k for k in agg.index if len(k) >= 4 and re.search(r"[^\W\d_]", k)]
    if len(keys) < 2:
        return []
    m = process.cdist(keys, keys, scorer=fuzz.ratio, dtype=np.uint8, workers=-1, score_cutoff=min_score)
    r, c = np.nonzero(np.triu(m, 1) >= min_score)
    pairs = []
    for i, j in zip(r, c):
        k1, k2 = keys[i], keys[j]
        if re.sub(r"\D", "", k1) != re.sub(r"\D", "", k2):   # "Room 101" vs "Room 102" are different
            continue
        a, b = (k1, k2) if agg.n[k1] >= agg.n[k2] else (k2, k1)
        pairs.append({"Value": agg.v[a], "Count": int(agg.n[a]), "Similar value": agg.v[b],
                      "Count ": int(agg.n[b]), "Similarity %": int(m[i, j])})
    pairs.sort(key=lambda p: -p["Similarity %"])
    return pairs[:limit]


# =============================================================================
# 4. DATA QUALITY CHECK (reports only - changes nothing)
# =============================================================================
def check_quality(df: pd.DataFrame, table: str = "A", cfg: dict | None = None) -> dict:
    """Return {'overview','issues','columns','similar','duplicates'} DataFrames."""
    cfg = {**DEFAULT_CLEANING, **(cfg or {})}
    markers = {str(m).strip().lower() for m in cfg["empty_markers"]}
    day_first, decimal = cfg["day_first"], cfg["decimal"]
    today = pd.Timestamp.today().normalize()
    n = len(df)
    issues, col_rows, similar = [], [], []

    def issue(col, level, problem, cells, examples, fix):
        issues.append({"Table": table, "Column": col, "Level": level, "Problem": problem,
                       "Cells": int(cells), "Examples": examples, "How to fix": fix})

    for c in df.columns:
        s = df[c]
        start = len(issues)
        strm = _str_mask(s)
        sv = s[strm].pipe(as_text)
        low = sv.str.strip().str.lower()
        marker_mask = pd.Series(False, index=s.index)
        marker_mask.loc[strm] = low.isin(markers).to_numpy()
        real = s[s.notna() & ~marker_mask]
        empty = int(s.isna().sum() + marker_mask.sum())
        kind = detect_kind(real, day_first, decimal)
        rs = real[_str_mask(real)].pipe(as_text)
        distinct = int(norm_text(real).nunique()) if len(real) else 0

        if n and empty == n:
            issue(c, "Check", "Column is completely empty", empty, "",
                  "Remove the column, or check the export/source")
            col_rows.append({"Table": table, "Column": c, "Type found": "empty", "Filled": 0,
                             "Empty": empty, "Empty %": 100.0, "Distinct": 0, "Examples": "",
                             "Min": "", "Max": "", "Issues": len(issues) - start})
            continue

        if empty:
            issue(c, "Info" if empty / n < 0.05 else "Check", f"Empty cells ({empty / n:.0%})", empty, "",
                  "Normal if optional. If it must be filled: RULES {'required': True}")
        ph = sv[low.isin(markers - {""})]
        if len(ph):
            issue(c, "Fix", "Placeholder text that means 'empty'", len(ph), _examples(ph),
                  "EMPTY_MARKERS (on by default when cleaning)")
        blank = sv[low == ""]
        if len(blank) and "" in markers:
            issue(c, "Fix", "Cells that contain only spaces", len(blank), "",
                  "EMPTY_MARKERS / TRIM_SPACES")

        if len(rs):
            tidy = tidy_spaces(rs)
            sp = rs[tidy != rs]
            if len(sp):
                issue(c, "Fix", "Extra spaces or invisible characters", len(sp),
                      _examples(repr(x) for x in sp), "TRIM_SPACES")
            key = tidy.str.lower()
            g = pd.DataFrame({"k": key, "v": rs}).drop_duplicates()
            multi = g.groupby("k")["v"].agg(list)
            multi = multi[multi.map(len) > 1]
            if len(multi):
                issue(c, "Fix", f"{len(multi)} value(s) written in different ways (upper/lower case, spaces)",
                      int(key.isin(multi.index).sum()),
                      " / ".join(repr(v) for v in multi.iloc[0][:4]), "UNIFY_SPELLING_VARIANTS")

        if kind == "number_text":
            issue(c, "Fix", "Numbers stored as text (can't be summed or compared correctly)",
                  len(rs), _examples(rs), "TO_NUMBER ('auto' handles it)")
            bad = real[parse_numbers(real, decimal)[1]]
            if len(bad):
                issue(c, "Check", "Values that are not valid numbers", len(bad), _examples(bad),
                      "Correct them in the source - cleaning leaves them empty and lists them")
        if kind == "date_text":
            issue(c, "Fix", "Dates stored as text", len(rs), _examples(rs), "TO_DATE ('auto' handles it)")
            shapes = (rs.str.translate(ARABIC_DIGITS).str.strip().str.replace(r"\d", "9", regex=True)
                        .str.replace(r"(?i)[a-z]+", "Mon", regex=True))
            if shapes.nunique() > 1:
                issue(c, "Check", f"Mixed date formats ({shapes.nunique()} different)", len(rs),
                      _examples(shapes.value_counts().index), "TO_DATE handles mixed formats")
            parts = rs.str.translate(ARABIC_DIGITS).str.extract(r"^\s*(\d{1,2})[/.-](\d{1,2})[/.-]\d{2,4}")
            p1, p2 = pd.to_numeric(parts[0], errors="coerce"), pd.to_numeric(parts[1], errors="coerce")
            amb = rs[((p1 <= 12) & (p2 <= 12) & (p1 != p2)).fillna(False).to_numpy()]
            if len(amb):
                issue(c, "Check", "Dates that can be read 2 ways (day/month or month/day)", len(amb),
                      _examples(amb), f"DAY_FIRST is {day_first} - make sure that is right for this file")
            bad = real[parse_dates(real, day_first)[1]]
            if len(bad):
                issue(c, "Check", "Values that are not valid dates", len(bad), _examples(bad),
                      "Correct them in the source - cleaning leaves them empty and lists them")

        if kind == "text":
            other = "," if decimal == "." else "."
            if detect_kind(real, day_first, other) == "number_text":
                issue(c, "Fix", f"Numbers written with '{other}' as the decimal separator", len(rs),
                      _examples(rs), f'DECIMAL_SEPARATOR = "{other}"  (then TO_NUMBER converts them)')

        vmin = vmax = ""
        if kind in ("number", "number_text"):
            vals = parse_numbers(real, decimal)[0].dropna()
            if len(vals):
                vmin, vmax = vals.min(), vals.max()
                q1, q3 = vals.quantile([0.25, 0.75])
                iqr = q3 - q1
                if iqr > 0 and len(vals) >= 8:
                    out = vals[(vals < q1 - 3 * iqr) | (vals > q3 + 3 * iqr)]
                    if len(out):
                        issue(c, "Check", "Unusually large or small values", len(out),
                              _examples(out.sort_values(key=abs, ascending=False)),
                              "Check them. To set limits: RULES {'min': ..., 'max': ...}")
                neg = vals[vals < 0]
                if 0 < len(neg) <= max(1, 0.05 * len(vals)) and len(vals) >= 5:
                    issue(c, "Check", "A few negative values (most are positive)", len(neg),
                          _examples(neg), "If not allowed: RULES {'min': 0}")
        if kind in ("date", "date_text"):
            dvals = parse_dates(real, day_first)[0].dropna()
            if len(dvals):
                vmin, vmax = dvals.min().date(), dvals.max().date()
                fut = dvals[dvals > today]
                if len(fut):
                    issue(c, "Check", "Dates in the future", len(fut), _examples(fut.dt.date),
                          "If not allowed: RULES {'not_future': True}")
                old = dvals[dvals.dt.year < 1900]
                if len(old):
                    issue(c, "Check", "Very old dates (before 1900) - often typing errors", len(old),
                          _examples(old.dt.date), "Correct them in the source")

        if pd.api.types.is_object_dtype(s) and kind not in ("number_text", "date_text"):
            types = real.map(lambda x: "number" if isinstance(x, (int, float, np.number))
                             and not isinstance(x, bool) else ("text" if isinstance(x, str)
                                                               else type(x).__name__))
            vc = types.value_counts()
            if len(vc) > 1:
                issue(c, "Check", "Mixed value types in one column", len(real),
                      ", ".join(f"{k}: {v}" for k, v in vc.items()),
                      "Usually from manual typing - check the column")

        if re.search(ID_HINT, str(c)) and len(real) > 1:
            dup_ids = real[norm_text(real, True).duplicated(keep=False)]
            if len(dup_ids):
                issue(c, "Check", "Repeated values in an ID/number column", len(dup_ids),
                      _examples(dup_ids), "OK if one ID can have several rows. Otherwise: RULES {'unique': True}")
        if distinct == 1 and len(real) > 1:
            issue(c, "Info", "Same value in every filled row", len(real), _examples(real), "Nothing to do")
        mh = re.match(r"^(.*)\.(\d+)$", str(c))
        if mh and mh.group(1) in df.columns:
            issue(c, "Check", f"Looks like a second '{mh.group(1)}' column (duplicate header)", 0, "",
                  "Rename one of them in the source file")

        if kind == "text" and 1 < distinct <= 20000:
            sims = similar_values(rs)
            if sims:
                for p in sims:
                    similar.append({"Table": table, "Column": c, **p})
                issue(c, "Check", f"Possible typos: {len(sims)} pair(s) of very similar values",
                      sum(p["Count "] for p in sims),
                      " | ".join(f"'{p['Similar value']}' ~ '{p['Value']}'" for p in sims[:3]),
                      "If they are the same thing: REPLACE_VALUES. See 'Similar spellings' sheet")

        type_label = TYPE_LABELS[kind]
        if distinct == len(real) and len(real) > 1 and (kind in ("text", "code") or re.search(ID_HINT, str(c))):
            type_label += " - all unique (ID?)"
        col_rows.append({"Table": table, "Column": c, "Type found": type_label, "Filled": len(real),
                         "Empty": empty, "Empty %": round(100 * empty / n, 1) if n else 0,
                         "Distinct": distinct, "Examples": _examples(real.head(200)),
                         "Min": vmin, "Max": vmax, "Issues": len(issues) - start})

    # ---- whole-table checks ----
    exact = int(df.duplicated().sum())
    normed = df.pipe(as_text).apply(lambda col: col.str.replace(r"\s+", " ", regex=True).str.strip().str.lower())
    hidden = int(normed.duplicated().sum()) - exact
    if exact:
        issue("(whole row)", "Fix", "Identical duplicate rows", exact, "", "REMOVE_DUPLICATE_ROWS")
    if hidden > 0:
        issue("(whole row)", "Fix", "Rows that are duplicates except for spaces/upper-lower case",
              hidden, "", "Cleaning makes them identical, then REMOVE_DUPLICATE_ROWS removes them")
    dup_mask = normed.duplicated(keep=False)
    dups = df[dup_mask.to_numpy()].copy()
    dups.insert(0, "Row in file", file_row(dups.index))
    dups.insert(0, "Table", table)
    dups = dups.head(5000)

    issues_df = pd.DataFrame(issues, columns=["Table", "Column", "Level", "Problem", "Cells",
                                              "Examples", "How to fix"])
    order = {"Fix": 0, "Check": 1, "Info": 2}
    issues_df = issues_df.sort_values("Level", key=lambda x: x.map(order), kind="stable").reset_index(drop=True)
    total_cells = n * len(df.columns)
    empty_cells = sum(r["Empty"] for r in col_rows)
    overview = pd.DataFrame([
        ("Rows", n), ("Columns", len(df.columns)),
        ("Empty cells", f"{empty_cells} ({empty_cells / total_cells:.1%})" if total_cells else 0),
        ("Identical duplicate rows", exact), ("Hidden duplicates (spaces/case)", max(hidden, 0)),
        ("Problems to fix", int((issues_df["Level"] == "Fix").sum())),
        ("Things to check", int((issues_df["Level"] == "Check").sum())),
    ], columns=["Item", "Value"])
    overview.insert(0, "Table", table)
    sim_df = pd.DataFrame(similar, columns=["Table", "Column", "Value", "Count", "Similar value",
                                            "Count ", "Similarity %"])
    return {"overview": overview, "issues": issues_df, "columns": pd.DataFrame(col_rows),
            "similar": sim_df, "duplicates": dups}


# =============================================================================
# 5. CLEANING + RULES
# =============================================================================
def validate_cleaning(cfg: dict, tables: dict):
    """Check cleaning settings against the columns of the loaded tables."""
    all_cols = []
    for df in tables.values():
        all_cols += [str(c) for c in df.columns if str(c) not in all_cols]

    def need(cols, setting):
        bad = [c for c in cols if c not in all_cols]
        if bad:
            lines = [f'  - "{c}" is not a column in any table{_suggest(c, all_cols)}' for c in bad]
            raise SettingsError(f"Problem in setting {setting}:\n" + "\n".join(lines)
                                + f"\n  Columns available: {all_cols}")

    for key, setting in (("to_number", "TO_NUMBER"), ("to_date", "TO_DATE")):
        v = cfg[key]
        if isinstance(v, (list, tuple)):
            need(v, setting)
        elif v != "auto":
            raise SettingsError(f'{setting} must be "auto" or a list like ["Column"] ([] = off).')
    if cfg["decimal"] not in (".", ","):
        raise SettingsError('DECIMAL_SEPARATOR must be "." or ",".')
    if not isinstance(cfg["fix_case"], dict):
        raise SettingsError('FIX_CASE must look like {"City": "title"} ({} = off).')
    need(cfg["fix_case"], "FIX_CASE")
    for col, mode in cfg["fix_case"].items():
        if mode not in CASE_MODES:
            raise SettingsError(f'FIX_CASE["{col}"] = "{mode}" is not valid. Use one of: {sorted(CASE_MODES)}')
    if not isinstance(cfg["replace"], dict):
        raise SettingsError('REPLACE_VALUES must look like {"City": {"old": "New"}} ({} = off).')
    need(cfg["replace"], "REPLACE_VALUES")
    for col, mp in cfg["replace"].items():
        if not isinstance(mp, dict):
            raise SettingsError(f'REPLACE_VALUES["{col}"] must be a dictionary like {{"old": "New"}}.')
    merge = cfg.get("merge") or {}
    if merge:
        keys = merge.get("key")
        if not isinstance(keys, (list, tuple)) or not keys:
            raise SettingsError('MERGE_DUPLICATE_KEYS needs a "key" list, e.g. {"key": ["Invoice No"]}.')
        need(keys, 'MERGE_DUPLICATE_KEYS["key"]')
        how = merge.get("how", {})
        need(how, 'MERGE_DUPLICATE_KEYS["how"]')
        for col, h in how.items():
            if h not in MERGE_HOW:
                raise SettingsError(f'MERGE_DUPLICATE_KEYS how "{h}" for "{col}" is not valid. '
                                    f"Use one of: {sorted(MERGE_HOW)}")
    if not isinstance(cfg["rules"], (list, tuple)):
        raise SettingsError("RULES must be a list of rules: [ {...}, {...} ]  ([] = off).")
    for n, r in enumerate(cfg["rules"], 1):
        if not isinstance(r, dict) or "column" not in r:
            raise SettingsError(f'RULES item {n} needs a "column", e.g. {{"column": "Amount", "min": 0}}.')
        need([r["column"]], f"RULES item {n}")
        checks = set(r) - {"column"}
        unknown = checks - RULE_CHECKS
        if unknown or not checks:
            raise SettingsError(f"RULES item {n}: unknown or missing check {sorted(unknown) or ''}. "
                                f"Use: {sorted(RULE_CHECKS)}")
        if "not_after" in r:
            need([r["not_after"]], f'RULES item {n} "not_after"')
        if "pattern" in r:
            try:
                re.compile(r["pattern"])
            except re.error as e:
                raise SettingsError(f"RULES item {n}: pattern is not valid ({e}).")


def evaluate_rules(df: pd.DataFrame, rules, table="A", day_first=True) -> list:
    """Return a list of problem rows. Nothing is deleted."""
    problems = []
    today = pd.Timestamp.today().normalize()

    def add(col, mask, text, values):
        mask = pd.Series(mask, index=df.index).fillna(False).astype(bool)
        for i in df.index[mask.to_numpy()][:20000]:
            problems.append({"Table": table, "Row in file": int(file_row(i)), "Column": col,
                             "Value": _show(values.loc[i]), "Problem": text})

    for r in rules:
        col = r["column"]
        if col not in df.columns:
            continue
        s = df[col]
        is_dt = pd.api.types.is_datetime64_any_dtype(s)
        if r.get("required"):
            empty = s.isna() | (_str_mask(s) & (s.pipe(as_text).str.strip() == ""))
            add(col, empty, "Empty, but this column is required", s)
        if r.get("unique"):
            add(col, s.notna() & norm_text(s, True).duplicated(keep=False),
                "Repeated value, but this column must be unique", s)
        for key, op, word in (("min", "lt", "at least"), ("max", "gt", "at most")):
            if key in r:
                if is_dt:
                    vals, bound = s, pd.Timestamp(r[key])
                else:
                    vals, bound = parse_numbers(s)[0], float(r[key])
                add(col, getattr(vals, op)(bound), f"Must be {word} {r[key]}", s)
        if "allowed" in r:
            allowed = {str(x).strip().lower() for x in r["allowed"]}
            add(col, s.notna() & ~norm_text(s, True).isin(allowed),
                f"Not an allowed value ({', '.join(map(str, r['allowed']))})", s)
        if "pattern" in r:
            ok = s.pipe(as_text).str.fullmatch(r["pattern"]).fillna(False).astype(bool)
            add(col, s.notna() & ~ok, f"Does not match the required pattern {r['pattern']}", s)
        if r.get("not_future"):
            dates = s if is_dt else parse_dates(s, day_first)[0]
            add(col, dates > today, "Date is in the future", s)
        if "not_after" in r and r["not_after"] in df.columns:
            o = df[r["not_after"]]
            if is_dt or pd.api.types.is_datetime64_any_dtype(o):
                left = s if is_dt else parse_dates(s, day_first)[0]
                right = o if pd.api.types.is_datetime64_any_dtype(o) else parse_dates(o, day_first)[0]
            else:
                left, right = parse_numbers(s)[0], parse_numbers(o)[0]
            add(col, left > right, f"Is after/bigger than '{r['not_after']}'", s)
    return problems


def clean_table(df: pd.DataFrame, cfg: dict | None = None, table: str = "A"):
    """Clean one table. Returns (cleaned DataFrame, log rows, problem rows)."""
    cfg = {**DEFAULT_CLEANING, **(cfg or {})}
    df = df.copy()
    log, problems = [], []

    def note(step, col, before, after, extra=""):
        b, a = before.astype(object), after.astype(object)
        changed = ~((b == a).fillna(False).astype(bool) | (b.isna() & a.isna()))
        k = int(changed.sum())
        if k:
            i = changed.idxmax()
            log.append({"Table": table, "Step": step, "Column": col, "Cells changed": k,
                        "Example before": repr(b[i]) if isinstance(b[i], str) else b[i],
                        "Example after": _show(a[i]), "Note": extra})

    def text_cols():
        return [c for c in df.columns if is_text_col(df[c])]

    # 1) placeholders -> empty
    markers = {str(m).strip().lower() for m in cfg["empty_markers"]}
    if markers:
        for c in text_cols():
            s = df[c]
            strm = _str_mask(s)
            hit = pd.Series(False, index=s.index)
            hit.loc[strm] = s[strm].pipe(as_text).str.strip().str.lower().isin(markers).to_numpy()
            if hit.any():
                new = s.astype(object).mask(hit, np.nan)
                note("Placeholder -> empty", c, s, new)
                df[c] = new

    # 2) spaces / invisible characters
    if cfg["trim"]:
        for c in text_cols():
            new = map_text(df[c], tidy_spaces)
            note("Trimmed spaces / invisible characters", c, df[c], new)
            df[c] = new

    # 3) known replacements (ignores case and extra spaces)
    for c, mapping in (cfg["replace"] or {}).items():
        if c not in df.columns:
            continue
        lookup = {re.sub(r"\s+", " ", str(k)).strip().lower(): v for k, v in mapping.items()}
        s = df[c]
        strm = _str_mask(s)
        key = s[strm].pipe(as_text).str.replace(r"\s+", " ", regex=True).str.strip().str.lower()
        hit = key[key.isin(lookup)]
        if len(hit):
            new = s.astype(object).copy()
            new.loc[hit.index] = hit.map(lookup).to_numpy()
            note("Replaced values (REPLACE_VALUES)", c, s, new)
            df[c] = new

    # 4) forced case
    for c, mode in (cfg["fix_case"] or {}).items():
        if c in df.columns:
            new = map_text(df[c], lambda t, m=mode: getattr(t.str, m)())
            note(f"Case -> {mode}", c, df[c], new)
            df[c] = new

    # 5) spelling variants (case/spaces only) -> most common spelling
    if cfg["unify"]:
        cols = cfg["unify"] if isinstance(cfg["unify"], (list, tuple)) else text_cols()
        for c in cols:
            if c not in df.columns:
                continue
            s = df[c]
            strm = _str_mask(s)
            if not strm.any():
                continue
            sv = s[strm].pipe(as_text)
            key = sv.str.replace(r"\s+", " ", regex=True).str.strip().str.lower()
            counts = pd.DataFrame({"k": key, "v": sv}).value_counts().reset_index()
            canon = counts.drop_duplicates("k").set_index("k")["v"]
            new = s.astype(object).copy()
            new.loc[strm] = key.map(canon).to_numpy()
            note("Unified spelling variants (case/spaces)", c, s, new)
            df[c] = new

    # 6) dates (before numbers, so '2026-03-01' is never read as a number)
    if cfg["to_date"] == "auto":
        date_cols = [c for c in text_cols() if detect_kind(df[c], cfg["day_first"], cfg["decimal"]) == "date_text"]
    else:
        date_cols = [c for c in cfg["to_date"] if c in df.columns]
    for c in date_cols:
        dates, failed = parse_dates(df[c], cfg["day_first"])
        note("Text -> date", c, df[c], dates,
             f"{int(failed.sum())} value(s) could not be read - see Data problems" if failed.any() else "")
        for i in df.index[failed.to_numpy()]:
            problems.append({"Table": table, "Row in file": int(file_row(i)), "Column": c,
                             "Value": df.at[i, c], "Problem": "Not a valid date - left empty"})
        df[c] = dates

    # 7) numbers
    if cfg["to_number"] == "auto":
        num_cols = [c for c in text_cols() if detect_kind(df[c], cfg["day_first"], cfg["decimal"]) == "number_text"]
    else:
        num_cols = [c for c in cfg["to_number"] if c in df.columns and c not in date_cols]
    for c in num_cols:
        nums, failed = parse_numbers(df[c], cfg["decimal"])
        whole = nums.dropna()
        if len(whole) and (whole % 1 == 0).all() and whole.abs().max() < 2**53:
            nums = nums.astype("Int64")
        note("Text -> number", c, df[c], nums,
             f"{int(failed.sum())} value(s) could not be read - see Data problems" if failed.any() else "")
        for i in df.index[failed.to_numpy()]:
            problems.append({"Table": table, "Row in file": int(file_row(i)), "Column": c,
                             "Value": df.at[i, c], "Problem": "Not a valid number - left empty"})
        df[c] = nums

    # 8) rules (flag only)
    problems += evaluate_rules(df, cfg["rules"] or [], table, cfg["day_first"])

    # 9) identical duplicates
    if cfg["remove_dups"]:
        dup = df.duplicated(keep="first")
        if dup.any():
            rows = file_row(df.index[dup.to_numpy()])
            log.append({"Table": table, "Step": "Removed identical duplicate rows", "Column": "(whole row)",
                        "Cells changed": int(dup.sum()),
                        "Example before": "rows " + ", ".join(map(str, rows[:15])) + (" ..." if len(rows) > 15 else ""),
                        "Example after": "removed", "Note": "Kept the first copy of each"})
            df = df[~dup.to_numpy()]

    # 10) merge rows sharing a key
    merge = cfg.get("merge") or {}
    if merge and all(k in df.columns for k in merge["key"]):
        keys, how = list(merge["key"]), merge.get("how", {})

        def agg_for(col):
            h = how.get(col, "first")
            if h == "join":
                return lambda x: ", ".join(dict.fromkeys(str(v) for v in x.dropna())) or np.nan
            return h

        before = len(df)
        tmp = df.rename_axis("_orig_row").reset_index()
        aggs = {"_orig_row": "first", **{c: agg_for(c) for c in df.columns if c not in keys}}
        merged = tmp.groupby(keys, dropna=False, sort=False).agg(aggs).reset_index()
        merged = merged.set_index("_orig_row")[list(df.columns)]
        merged.index.name = None
        if len(merged) < before:
            log.append({"Table": table, "Step": "Merged rows with the same key", "Column": ", ".join(keys),
                        "Cells changed": before - len(merged), "Example before": f"{before} rows",
                        "Example after": f"{len(merged)} rows",
                        "Note": "; ".join(f"{k}: {v}" for k, v in how.items()) or "first value kept"})
        df = merged
    return df, log, problems


def logs_to_frames(log, problems):
    log_df = pd.DataFrame(log, columns=["Table", "Step", "Column", "Cells changed",
                                        "Example before", "Example after", "Note"])
    prob_df = pd.DataFrame(problems, columns=["Table", "Row in file", "Column", "Value", "Problem"])
    return log_df, prob_df


# =============================================================================
# 6. PIVOT
# =============================================================================
def make_pivot(df, rows, columns, values, agg, totals=True, fill0=True) -> pd.DataFrame:
    data = df.copy()
    group_cols = rows + columns
    data[group_cols] = data[group_cols].astype(object).where(data[group_cols].notna(), "(blank)")

    if not values:                                   # nothing to calculate -> count rows
        data["Row count"] = 1
        values, agg = ["Row count"], "sum"
    elif agg in NUMERIC_AGGS:
        for v in values:
            data[v] = pd.to_numeric(data[v], errors="coerce").astype(float)

    swap = not rows                                  # only columns chosen -> build then flip
    pt = pd.pivot_table(data, index=columns if swap else rows,
                        columns=None if swap else (columns or None),
                        values=values, aggfunc=agg, margins=totals, margins_name="Total",
                        fill_value=0 if fill0 else None)
    if swap:
        pt = pt.T
    if isinstance(pt.columns, pd.MultiIndex) and len(values) == 1:
        pt.columns = pt.columns.droplevel(0)
    pt = pt.reset_index()
    if swap:
        pt = pt.rename(columns={"index": "Value"})
    return flatten_columns(pt)


# =============================================================================
# 7. EXACT TABLE COMPARISON
# =============================================================================
def compare_schema(a: pd.DataFrame, b: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for c in dict.fromkeys(list(a.columns) + list(b.columns)):
        in_a, in_b = c in a.columns, c in b.columns
        rows.append({"Column": c,
                     "Status": "In both" if in_a and in_b else ("Only in A" if in_a else "Only in B"),
                     "Type in A": str(a[c].dtype) if in_a else "",
                     "Type in B": str(b[c].dtype) if in_b else ""})
    return pd.DataFrame(rows)


def compare_tables(a, b, keys, cols, tol=0.0, ignore_case=False) -> dict:
    """Match rows by key columns (same names in both tables) and report differences."""
    cols = [c for c in cols if c not in keys]
    a = a[keys + cols].copy()
    b = b[keys + cols].copy()
    for k in keys:                                   # 1001 / 1001.0 / " 1001" all match
        a[k] = norm_text(a[k], ignore_case)
        b[k] = norm_text(b[k], ignore_case)

    dup_a, dup_b = int(a.duplicated(keys).sum()), int(b.duplicated(keys).sum())
    m = a.merge(b, on=keys, how="outer", suffixes=(" (A)", " (B)"), indicator=True)

    only_a = m[m["_merge"] == "left_only"][keys + [f"{c} (A)" for c in cols]]
    only_a.columns = keys + cols
    only_b = m[m["_merge"] == "right_only"][keys + [f"{c} (B)" for c in cols]]
    only_b.columns = keys + cols
    both = m[m["_merge"] == "both"].reset_index(drop=True)

    parts, any_diff = [], np.zeros(len(both), dtype=bool)
    for c in cols:
        eq = values_equal(both[f"{c} (A)"], both[f"{c} (B)"], tol, ignore_case)
        any_diff |= ~eq
        d = both.loc[~eq, keys + [f"{c} (A)", f"{c} (B)"]].copy()
        if d.empty:
            continue
        d.columns = keys + ["Value A", "Value B"]
        d.insert(len(keys), "Column", c)
        parts.append(d)

    diffs = (pd.concat(parts, ignore_index=True) if parts
             else pd.DataFrame(columns=keys + ["Column", "Value A", "Value B"]))
    if not diffs.empty:
        diffs["Difference (B - A)"] = (pd.to_numeric(diffs["Value B"], errors="coerce")
                                       - pd.to_numeric(diffs["Value A"], errors="coerce"))
        diffs = diffs.sort_values(keys).reset_index(drop=True)

    return {"only_a": only_a.reset_index(drop=True), "only_b": only_b.reset_index(drop=True),
            "diffs": diffs, "matched": len(both), "matched_changed": int(any_diff.sum()),
            "dup_a": dup_a, "dup_b": dup_b}


def exact_compare(a, b, key_a, key_b, compare_cols, tol=0.0, ignore_case=False) -> dict:
    """Like compare_tables, but key/compare columns may have different names in A and B."""
    compare_cols = {ca: cb for ca, cb in compare_cols.items() if ca not in key_a}
    a2 = a[key_a + list(compare_cols)].copy()
    b2 = b[key_b + list(compare_cols.values())].copy()
    b2.columns = key_a + list(compare_cols)          # rename B's columns to A's names
    return compare_tables(a2, b2, key_a, list(compare_cols), tol, ignore_case)


# =============================================================================
# 8. FUZZY TABLE COMPARISON
# =============================================================================
def get_scorer(method: str):
    if fuzz is None:
        raise SettingsError("Fuzzy matching needs rapidfuzz. Install it with:  pip install rapidfuzz")
    if method not in MATCH_METHODS:
        raise SettingsError(f'MATCH_METHOD "{method}" is not valid. Use one of: {list(MATCH_METHODS)}')
    return getattr(fuzz, MATCH_METHODS[method][0])


def clean_for_matching(s: pd.Series, ignore_words=()) -> pd.Series:
    """lower case, remove punctuation and filler words, collapse spaces."""
    s = norm_text(s, True)
    s = s.str.translate(ARABIC_DIGITS).str.replace(r"[^\w\s]", " ", regex=True)   # \w keeps Arabic/Kurdish
    words = [w.strip().lower() for w in ignore_words if str(w).strip()]
    if words:
        s = s.str.replace(r"\b(" + "|".join(map(re.escape, words)) + r")\b", " ", regex=True)
    return s.str.replace(r"\s+", " ", regex=True).str.strip()


def _match_text(df, cols, ignore_words):
    out = clean_for_matching(df[cols[0]], ignore_words)
    for c in cols[1:]:
        out = out + " " + clean_for_matching(df[c], ignore_words)
    return out.str.strip().to_numpy(dtype=object)


def fuzzy_compare(a, b, key_a, key_b, compare_cols=None, threshold=85, method="token_sort",
                  must_match=None, ignore_words=(), review_below=95, tol=0.0,
                  ignore_case=True) -> dict:
    """
    Pair rows of A and B whose key text is similar (score 0-100).
    Each row is used at most once; the best-scoring pairs are chosen first.
    """
    scorer = get_scorer(method)
    must_match = must_match or {}
    rows_a, rows_b = a.index.to_numpy(), b.index.to_numpy()     # original file rows
    a = a.reset_index(drop=True)
    b = b.reset_index(drop=True)
    if not compare_cols:
        skip = set(key_a) | set(key_b) | set(must_match) | set(must_match.values())
        compare_cols = {c: c for c in a.columns if c in b.columns and c not in skip}

    ta, tb = _match_text(a, key_a, ignore_words), _match_text(b, key_b, ignore_words)

    ga = np.full(len(a), "", dtype=object)
    gb = np.full(len(b), "", dtype=object)
    for ca, cb in must_match.items():
        ga = ga + "|" + norm_text(a[ca], ignore_case).to_numpy(dtype=object)
        gb = gb + "|" + norm_text(b[cb], ignore_case).to_numpy(dtype=object)

    best_a, best_a_idx = np.zeros(len(a), dtype=int), np.full(len(a), -1)
    best_b, best_b_idx = np.zeros(len(b), dtype=int), np.full(len(b), -1)
    cand_s, cand_i, cand_j = [], [], []

    for g in pd.unique(ga):
        ia = np.where((ga == g) & (ta != ""))[0]
        ib = np.where((gb == g) & (tb != ""))[0]
        if len(ia) == 0 or len(ib) == 0:
            continue
        m = process.cdist(ta[ia].tolist(), tb[ib].tolist(), scorer=scorer, dtype=np.uint8, workers=-1)
        # second check with spaces removed: "Silver Line" = "Silverline"
        m2 = process.cdist([x.replace(" ", "") for x in ta[ia]], [x.replace(" ", "") for x in tb[ib]],
                           scorer=fuzz.ratio, dtype=np.uint8, workers=-1)
        m = np.maximum(m, m2)
        ba = m.argmax(axis=1)
        best_a[ia], best_a_idx[ia] = m[np.arange(len(ia)), ba], ib[ba]
        bb = m.argmax(axis=0)
        best_b[ib], best_b_idx[ib] = m[bb, np.arange(len(ib))], ia[bb]
        r, c = np.nonzero(m >= threshold)
        cand_s.append(m[r, c].astype(int))
        cand_i.append(ia[r])
        cand_j.append(ib[c])

    used_a, used_b = np.zeros(len(a), dtype=bool), np.zeros(len(b), dtype=bool)
    pairs = []
    if cand_s:
        s, i, j = np.concatenate(cand_s), np.concatenate(cand_i), np.concatenate(cand_j)
        for k in np.lexsort((j, i, -s)):
            if not used_a[i[k]] and not used_b[j[k]]:
                used_a[i[k]] = used_b[j[k]] = True
                pairs.append((int(i[k]), int(j[k]), int(s[k])))
    pairs.sort()
    pi = np.array([p[0] for p in pairs], dtype=int)
    pj = np.array([p[1] for p in pairs], dtype=int)
    ps = np.array([p[2] for p in pairs], dtype=int)

    matches = pd.DataFrame({"Row in A file": file_row(rows_a[pi]), "Row in B file": file_row(rows_b[pj])})
    for k in key_a:
        matches[f"{k} (A)"] = a[k].iloc[pi].to_numpy()
    for k in key_b:
        matches[f"{k} (B)"] = b[k].iloc[pj].to_numpy()
    if len(pi):
        raw_a = np.array([" ".join(x) for x in zip(*[norm_text(a[k], True).iloc[pi] for k in key_a])])
        raw_b = np.array([" ".join(x) for x in zip(*[norm_text(b[k], True).iloc[pj] for k in key_b])])
    else:
        raw_a = raw_b = np.array([])
    matches["Similarity %"] = ps
    matches["Match type"] = np.where(raw_a == raw_b, "Identical",
                                     np.where(ps == 100, "Same after cleaning", "Similar"))
    matches["Needs review"] = np.where(ps < review_below, "Yes", "")
    key_cols = [c for c in matches.columns if c.endswith(" (A)") or c.endswith(" (B)")]

    parts, n_diff = [], np.zeros(len(pairs), dtype=int)
    for ca, cb in compare_cols.items():
        va = a[ca].iloc[pi].reset_index(drop=True)
        vb = b[cb].iloc[pj].reset_index(drop=True)
        eq = values_equal(va, vb, tol, ignore_case)
        n_diff += ~eq
        if (~eq).any():
            d = matches.loc[~eq, key_cols + ["Similarity %"]].copy()
            d["Column"] = ca if ca == cb else f"{ca} <-> {cb}"
            d["Value A"] = va[~eq].to_numpy()
            d["Value B"] = vb[~eq].to_numpy()
            parts.append(d)
    if compare_cols:
        matches["Values"] = np.where(n_diff == 0, "All same", np.char.add(n_diff.astype(str), " different"))
    diffs = (pd.concat(parts, ignore_index=True) if parts
             else pd.DataFrame(columns=key_cols + ["Similarity %", "Column", "Value A", "Value B"]))
    if not diffs.empty:
        diffs["Difference (B - A)"] = (pd.to_numeric(diffs["Value B"], errors="coerce")
                                       - pd.to_numeric(diffs["Value A"], errors="coerce"))

    def unmatched(df, rows, used, best, best_idx, other, other_keys, label):
        idx = np.where(~used)[0]
        out = df.iloc[idx].copy()
        out.insert(0, f"Row in {label} file", file_row(rows[idx]))
        out[f"Closest in {'B' if label == 'A' else 'A'}"] = [
            " / ".join(str(other[k].iloc[best_idx[x]]) for k in other_keys) if best_idx[x] >= 0 else ""
            for x in idx]
        out["Closest %"] = [int(best[x]) if best_idx[x] >= 0 else None for x in idx]
        return out.reset_index(drop=True)

    only_a = unmatched(a, rows_a, used_a, best_a, best_a_idx, b, key_b, "A")
    only_b = unmatched(b, rows_b, used_b, best_b, best_b_idx, a, key_a, "B")

    return {"matches": matches, "diffs": diffs, "only_a": only_a, "only_b": only_b,
            "stats": {"Rows in A": len(a), "Rows in B": len(b), "Matched pairs": len(pairs),
                      "Matches needing review": int((ps < review_below).sum()),
                      "Matched with different values": int((n_diff > 0).sum()),
                      "Only in A (no match)": len(only_a), "Only in B (no match)": len(only_b)}}


# =============================================================================
# 9. COLUMNS AND ROWS
# =============================================================================
def column_profile(s: pd.Series) -> dict:
    p = {"Rows": len(s), "Empty": int(s.isna().sum()), "Distinct values": int(s.nunique())}
    if pd.api.types.is_datetime64_any_dtype(s):
        d = s.dropna()
        if len(d):
            p.update({"Earliest": d.min().date(), "Latest": d.max().date()})
        return p
    num = pd.to_numeric(s, errors="coerce").astype(float)
    is_numeric = s.notna().sum() > 0 and num.notna().sum() >= 0.8 * s.notna().sum()
    if is_numeric:
        p.update({"Sum": round(num.sum(), 4), "Average": round(num.mean(), 4),
                  "Min": num.min(), "Max": num.max()})
    else:
        mode = s.mode()
        p["Most common"] = mode.iloc[0] if not mode.empty else ""
    return p


def value_overlap(sa: pd.Series, sb: pd.Series, ignore_case=False):
    A = set(norm_text(sa.dropna(), ignore_case))
    B = set(norm_text(sb.dropna(), ignore_case))
    return sorted(A & B), sorted(A - B), sorted(B - A)


def rowwise_compare(sa, sb, name_a, name_b, tol=0.0, ignore_case=False) -> pd.DataFrame:
    n = min(len(sa), len(sb))
    sa, sb = sa.iloc[:n].reset_index(drop=True), sb.iloc[:n].reset_index(drop=True)
    out = pd.DataFrame({"Row": range(n), name_a: sa, name_b: sb})
    out["Same?"] = values_equal(sa, sb, tol, ignore_case)
    if not (pd.api.types.is_datetime64_any_dtype(sa) or pd.api.types.is_datetime64_any_dtype(sb)):
        diff = (pd.to_numeric(sb, errors="coerce").astype(float)
                - pd.to_numeric(sa, errors="coerce").astype(float))
        if diff.notna().any():
            out["Difference (B - A)"] = diff
    return out


def find_row(df: pd.DataFrame, by: str, value):
    """Return (row, message). by = 'row_number' or a column name."""
    if str(by).lower().replace(" ", "_") == "row_number":
        try:
            i = int(value)
        except (TypeError, ValueError):
            return None, "Row number must be a whole number."
        if not 0 <= i < len(df):
            return None, f"Row number must be between 0 and {len(df) - 1}."
        return df.iloc[i], ""
    if pd.api.types.is_datetime64_any_dtype(df[by]):
        target = _one_date(str(value), True)
        hits = df[df[by].dt.normalize() == (target.normalize() if pd.notna(target) else target)]
    else:
        target = str(value).strip().lower().removesuffix(".0")
        hits = df[norm_text(df[by], True) == target]
    if hits.empty:
        return None, f"No row where {by} = '{value}'."
    return hits.iloc[0], (f"{len(hits)} rows matched; using the first." if len(hits) > 1 else "")


def compare_two_rows(ra: pd.Series, rb: pd.Series, ignore_case=True, tol=0.0,
                     pairs: dict | None = None) -> pd.DataFrame:
    """Compare same-named columns, plus any extra {col_in_A: col_in_B} pairs."""
    mapping = {c: c for c in ra.index if c in rb.index}
    for ca, cb in (pairs or {}).items():
        if ca in ra.index and cb in rb.index:
            mapping[ca] = cb
    out = pd.DataFrame({"Column": [ca if ca == cb else f"{ca} <-> {cb}" for ca, cb in mapping.items()],
                        "Row A": [ra[ca] for ca in mapping],
                        "Row B": [rb[cb] for cb in mapping.values()]})
    out["Same?"] = values_equal(out["Row A"], out["Row B"], tol, ignore_case)
    return out


# =============================================================================
# 10. EXCEL OUTPUT
# =============================================================================
def write_excel(sheets: dict, target) -> None:
    """Write {sheet_name: DataFrame} to a path or BytesIO, with simple formatting."""
    from openpyxl.styles import Font, PatternFill
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="305496")
    fills = {"Yes": PatternFill("solid", fgColor="FFF2CC"),     # needs review
             "Fix": PatternFill("solid", fgColor="F8CBAD"),     # problem to fix
             "Check": PatternFill("solid", fgColor="FFF2CC")}   # worth checking

    with pd.ExcelWriter(target, engine="openpyxl") as writer:
        for name, df in sheets.items():
            sn = re.sub(r"[\[\]:*?/\\]", "_", str(name))[:31] or "Sheet"
            if len(df) > EXCEL_MAX_ROWS:
                df = df.head(EXCEL_MAX_ROWS)
            df.to_excel(writer, sheet_name=sn, index=False)
            ws = writer.sheets[sn]
            ws.freeze_panes = "A2"
            for cell in ws[1]:
                cell.font, cell.fill = header_font, header_fill
            for col in ws.iter_cols(max_row=min(ws.max_row, 1000)):
                longest = max((len(str(c.value)) for c in col if c.value is not None), default=8)
                ws.column_dimensions[col[0].column_letter].width = min(longest + 2, 60)
            if len(df) <= 200000:                     # show dates without 00:00:00
                for k, col_name in enumerate(df.columns):
                    col = df[col_name]
                    if pd.api.types.is_datetime64_any_dtype(col) and (col.dropna() == col.dropna().dt.normalize()).all():
                        for (cell,) in ws.iter_rows(min_row=2, min_col=k + 1, max_col=k + 1):
                            cell.number_format = "yyyy-mm-dd"
            for flag_col in ("Needs review", "Level"):
                if flag_col in df.columns and len(df) < 50000:
                    k = list(df.columns).index(flag_col)
                    for row in ws.iter_rows(min_row=2):
                        fill = fills.get(row[k].value)
                        if fill:
                            for c in row:
                                c.fill = fill


def to_excel_bytes(sheets: dict) -> bytes:
    buf = io.BytesIO()
    write_excel(sheets, buf)
    return buf.getvalue()


def save_table(df: pd.DataFrame, path: str):
    """Save a cleaned table as CSV (Excel-friendly UTF-8) or XLSX."""
    folder = os.path.dirname(path)
    if folder:
        os.makedirs(folder, exist_ok=True)
    if path.lower().endswith(".csv"):
        df.to_csv(path, index=False, encoding="utf-8-sig")
    else:
        write_excel({"Data": df}, path)


# =============================================================================
# 11. RUN FROM SETTINGS
# =============================================================================
def _need_list(value, setting):
    if not isinstance(value, (list, tuple)) or not value:
        raise SettingsError(f'{setting} must be a list with at least one column, e.g. ["Name"]')
    return list(value)


def cleaning_config_from_settings() -> dict:
    return {"empty_markers": list(EMPTY_MARKERS), "trim": TRIM_SPACES, "unify": UNIFY_SPELLING_VARIANTS,
            "to_number": TO_NUMBER, "decimal": DECIMAL_SEPARATOR, "to_date": TO_DATE,
            "day_first": DAY_FIRST, "fix_case": dict(FIX_CASE or {}), "replace": dict(REPLACE_VALUES or {}),
            "remove_dups": REMOVE_DUPLICATE_ROWS, "merge": dict(MERGE_DUPLICATE_KEYS or {}),
            "rules": list(RULES or [])}


def _cleaned_path(source: str, label: str) -> str:
    base, ext = os.path.splitext(os.path.basename(source))
    ext = ".csv" if ext.lower() in (".csv", ".txt", ".tsv") else ".xlsx"
    return os.path.join(CLEANED_FOLDER, f"{base}_cleaned{ext}")


def run_from_settings():
    """Returns (sheets dict, summary lines)."""
    task = str(TASK).strip().lower()
    valid = ["check_data", "clean_data", "fuzzy_compare", "compare_tables", "compare_columns",
             "compare_rows", "pivot"]
    if task not in valid:
        raise SettingsError(f'TASK "{TASK}" is not valid. Use one of: {valid}')

    single_ok = task in ("check_data", "clean_data")
    want_b = task != "pivot" and not (single_ok and not INCLUDE_TABLE_B)
    if USE_DEMO_DATA:
        a, b = demo_tables()
        name_a, name_b = "Demo: ERP suppliers", "Demo: supplier statement"
        b = b if want_b else None
    else:
        a = load_table(TABLE_A, "TABLE_A")
        b = None
        if want_b and not (single_ok and not str(TABLE_B.get("file", "")).strip()):
            b = load_table(TABLE_B, "TABLE_B")
        name_a, name_b = TABLE_A["file"], TABLE_B.get("file", "")

    info = [("Task", task), ("Table A", name_a)]
    if b is not None:
        info.append(("Table B", name_b))
    info.append(("Run at", datetime.now().strftime("%Y-%m-%d %H:%M")))
    tables = {"A": a} if b is None else {"A": a, "B": b}

    cfg = cleaning_config_from_settings()
    do_clean = APPLY_CLEANING or task == "clean_data"
    if do_clean:
        validate_cleaning(cfg, tables)

    def run_cleaning():
        out, logs, probs = {}, [], []
        for label, df in tables.items():
            cdf, lg, pr = clean_table(df, cfg, label)
            out[label] = cdf
            logs += lg
            probs += pr
        log_df, prob_df = logs_to_frames(logs, probs)
        return out, log_df, prob_df

    # ---------- CHECK DATA ----------
    if task == "check_data":
        parts = [check_quality(df, label, cfg) for label, df in tables.items()]
        cat = {k: pd.concat([p[k] for p in parts if not p[k].empty] or [parts[0][k]], ignore_index=True)
               for k in parts[0]}
        n_fix = int((cat["issues"]["Level"] == "Fix").sum())
        n_check = int((cat["issues"]["Level"] == "Check").sum())
        info += [("Problems to fix", n_fix), ("Things to check", n_check),
                 ("Possible typos (similar spellings)", len(cat["similar"]))]
        sheets = {"Summary": pd.DataFrame(info, columns=["Item", "Value"]), "Overview": cat["overview"],
                  "Issues": cat["issues"], "Columns": cat["columns"],
                  "Similar spellings": cat["similar"], "Duplicate rows": cat["duplicates"]}
        if APPLY_CLEANING:
            cleaned, log_df, prob_df = run_cleaning()
            after = pd.concat([check_quality(df, label, cfg)["issues"] for label, df in cleaned.items()],
                              ignore_index=True)
            info += [("Cleaning would change", f"{int(log_df['Cells changed'].sum())} cells (see Cleaning log)"),
                     ("Fix-level problems left after cleaning", int((after["Level"] == "Fix").sum())),
                     ("Data problems (rules / unreadable values)", len(prob_df))]
            sheets["Summary"] = pd.DataFrame(info, columns=["Item", "Value"])
            sheets.update({"Issues after cleaning": after, "Cleaning log": log_df, "Data problems": prob_df})
        return sheets, info

    # ---------- CLEAN DATA ----------
    if task == "clean_data":
        cleaned, log_df, prob_df = run_cleaning()
        info += [("Cells changed", int(log_df["Cells changed"].sum())),
                 ("Data problems (rules / unreadable values)", len(prob_df))]
        sheets = {"Summary": None, "Cleaning log": log_df, "Data problems": prob_df}
        for label, df in cleaned.items():
            src = (name_a if label == "A" else name_b)
            path = _cleaned_path("demo_table_" + label if USE_DEMO_DATA else src, label)
            save_table(df, path)
            info += [(f"Table {label} rows", f"{len(tables[label])} -> {len(df)}"),
                     (f"Cleaned file {label}", os.path.abspath(path))]
            sheets[f"Cleaned {label}"] = df
        sheets["Summary"] = pd.DataFrame(info, columns=["Item", "Value"])
        return sheets, info

    # ---------- every other task: optionally clean first ----------
    extra = {}
    if APPLY_CLEANING:
        cleaned, log_df, prob_df = run_cleaning()
        a, b = cleaned["A"], cleaned.get("B")
        info.append(("Cleaning", f"on - {int(log_df['Cells changed'].sum())} cells changed, "
                                 f"{len(prob_df)} data problems (see sheets)"))
        extra = {"Cleaning log": log_df, "Data problems": prob_df}
    else:
        info.append(("Cleaning", "off"))

    def finish(sheets):
        sheets = {"Summary": pd.DataFrame(info, columns=["Item", "Value"]), **sheets, **extra}
        return sheets, info

    # ---------- PIVOT ----------
    if task == "pivot":
        check_columns(a, PIVOT_ROWS + PIVOT_COLUMNS + PIVOT_VALUES, "TABLE_A",
                      "PIVOT_ROWS / PIVOT_COLUMNS / PIVOT_VALUES")
        if not PIVOT_ROWS and not PIVOT_COLUMNS:
            raise SettingsError("Put at least one column in PIVOT_ROWS or PIVOT_COLUMNS.")
        if PIVOT_CALC not in AGG_FUNCS.values():
            raise SettingsError(f'PIVOT_CALC "{PIVOT_CALC}" is not valid. Use one of: {list(AGG_FUNCS.values())}')
        result = make_pivot(a, list(PIVOT_ROWS), list(PIVOT_COLUMNS), list(PIVOT_VALUES),
                            PIVOT_CALC, PIVOT_TOTALS)
        info += [("Rows", ", ".join(PIVOT_ROWS)), ("Columns", ", ".join(PIVOT_COLUMNS)),
                 ("Values", ", ".join(PIVOT_VALUES) or "row count"), ("Calculation", PIVOT_CALC),
                 ("Result size", f"{len(result)} rows")]
        return finish({"Pivot": result})

    # ---------- TABLE COMPARISONS ----------
    if task in ("fuzzy_compare", "compare_tables"):
        key_a, key_b = _need_list(KEY_A, "KEY_A"), _need_list(KEY_B, "KEY_B")
        if len(key_a) != len(key_b):
            raise SettingsError("KEY_A and KEY_B must have the same number of columns.")
        check_columns(a, key_a, "TABLE_A", "KEY_A")
        check_columns(b, key_b, "TABLE_B", "KEY_B")
        cols = dict(COMPARE_COLUMNS or {})
        if not cols:
            cols = {c: c for c in a.columns if c in b.columns and c not in key_a + key_b}
        check_columns(a, list(cols), "TABLE_A", "COMPARE_COLUMNS (left side)")
        check_columns(b, list(cols.values()), "TABLE_B", "COMPARE_COLUMNS (right side)")
        info += [("Key A", ", ".join(key_a)), ("Key B", ", ".join(key_b)),
                 ("Compared columns", ", ".join(f"{x} <-> {y}" for x, y in cols.items()) or "none"),
                 ("Number tolerance", NUMBER_TOLERANCE)]

        if task == "fuzzy_compare":
            mm = dict(MUST_MATCH_EXACTLY or {})
            check_columns(a, list(mm), "TABLE_A", "MUST_MATCH_EXACTLY (left side)")
            check_columns(b, list(mm.values()), "TABLE_B", "MUST_MATCH_EXACTLY (right side)")
            r = fuzzy_compare(a, b, key_a, key_b, cols, SIMILARITY_THRESHOLD, MATCH_METHOD, mm,
                              IGNORE_WORDS, REVIEW_BELOW, NUMBER_TOLERANCE, IGNORE_CASE)
            info += [("Method", MATCH_METHOD), ("Threshold", SIMILARITY_THRESHOLD),
                     ("Review below", REVIEW_BELOW)] + list(r["stats"].items())
            return finish({"Matches": r["matches"], "Different values": r["diffs"],
                           "Only in A": r["only_a"], "Only in B": r["only_b"]})

        r = exact_compare(a, b, key_a, key_b, cols, NUMBER_TOLERANCE, IGNORE_CASE)
        info += [("Rows in A", len(a)), ("Rows in B", len(b)), ("Matched", r["matched"]),
                 ("Matched with different values", r["matched_changed"]),
                 ("Only in A", len(r["only_a"])), ("Only in B", len(r["only_b"]))]
        if r["dup_a"] or r["dup_b"]:
            info.append(("WARNING", f"Key not unique (duplicates A={r['dup_a']}, B={r['dup_b']}). "
                                    "Add another key column, or use MERGE_DUPLICATE_KEYS."))
        return finish({"Different values": r["diffs"], "Only in A": r["only_a"], "Only in B": r["only_b"]})

    # ---------- COLUMNS ----------
    if task == "compare_columns":
        check_columns(a, [COLUMN_A], "TABLE_A", "COLUMN_A")
        check_columns(b, [COLUMN_B], "TABLE_B", "COLUMN_B")
        sa, sb = a[COLUMN_A], b[COLUMN_B]
        la, lb = f"A: {COLUMN_A}", f"B: {COLUMN_B}"
        prof = pd.DataFrame({la: pd.Series(column_profile(sa)), lb: pd.Series(column_profile(sb))})
        prof = prof.reset_index().rename(columns={"index": "Metric"})
        both, only_a, only_b = value_overlap(sa, sb, IGNORE_CASE)
        info += [("Column A", COLUMN_A), ("Column B", COLUMN_B), ("Values in both", len(both)),
                 ("Values only in A", len(only_a)), ("Values only in B", len(only_b))]
        sheets = {"Column stats": prof, "Values in both": pd.DataFrame({"Value": both}),
                  "Values only in A": pd.DataFrame({"Value": only_a}),
                  "Values only in B": pd.DataFrame({"Value": only_b})}
        if COMPARE_BY_POSITION:
            rw = rowwise_compare(sa, sb, la, lb, NUMBER_TOLERANCE, IGNORE_CASE)
            info.append(("Rows that differ (by position)", int((~rw["Same?"]).sum())))
            sheets["Row by row"] = rw
        return finish(sheets)

    # ---------- ROWS ----------
    for cfg_row, df, setting in ((ROW_A, a, "ROW_A"), (ROW_B, b, "ROW_B")):
        by = cfg_row.get("find_by", "row_number")
        if str(by).lower().replace(" ", "_") != "row_number":
            check_columns(df, [by], "TABLE_A" if setting == "ROW_A" else "TABLE_B", f'{setting}["find_by"]')
    ra, msg_a = find_row(a, ROW_A.get("find_by", "row_number"), ROW_A.get("value"))
    rb, msg_b = find_row(b, ROW_B.get("find_by", "row_number"), ROW_B.get("value"))
    if ra is None:
        raise SettingsError(f"ROW_A: {msg_a}")
    if rb is None:
        raise SettingsError(f"ROW_B: {msg_b}")
    out = compare_two_rows(ra, rb, IGNORE_CASE, NUMBER_TOLERANCE, dict(COMPARE_COLUMNS or {}))
    info += [("Row A", f"{ROW_A.get('find_by')} = {ROW_A.get('value')}"),
             ("Row B", f"{ROW_B.get('find_by')} = {ROW_B.get('value')}"),
             ("Compared columns", len(out)), ("Columns that differ", int((~out["Same?"]).sum()))]
    for m in (msg_a, msg_b):
        if m:
            info.append(("Note", m))
    return finish({"Row comparison": out})


def main():
    try:
        sys.stdout.reconfigure(errors="replace")      # avoid crashes on old Windows consoles
    except Exception:
        pass
    try:
        sheets, info = run_from_settings()
        folder = os.path.dirname(OUTPUT_FILE)
        if folder:
            os.makedirs(folder, exist_ok=True)
        write_excel(sheets, OUTPUT_FILE)
    except SettingsError as e:
        print("\n❌ Please fix the SETTINGS:\n" + str(e) + "\n")
        sys.exit(1)
    except PermissionError:
        print(f'\n❌ Could not save "{OUTPUT_FILE}". Is it open in Excel? Close it and run again.\n')
        sys.exit(1)

    print("\n✅ Done\n" + "-" * 60)
    for k, v in info:
        print(f"{k:<42} {v}")
    print("-" * 60 + f"\nReport saved to: {os.path.abspath(OUTPUT_FILE)}\n")


if __name__ == "__main__":
    main()