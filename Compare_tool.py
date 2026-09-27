"""
compare_tool.py - Compare & pivot your data (settings version)
===============================================================

HOW TO USE
  1. Edit the SETTINGS section below. Look for the 👉 comments.
  2. Run:      python compare_tool.py
  3. Open the Excel report saved at OUTPUT_FILE.

First time? Leave USE_DEMO_DATA = True and just run it to see an example report.

Install once:  pip install pandas openpyxl rapidfuzz
               (+ xlrd for old .xls files, + pyarrow for .parquet files)

This file is also the "engine" for data_tool.py (the click version),
so keep both files in the same folder.
"""

# ╔══════════════════════════════════════════════════════════════════════════╗
# ║                  ✏️  SETTINGS - EDIT ONLY THIS PART                       ║
# ╚══════════════════════════════════════════════════════════════════════════╝

# 👉 STEP 0 - Test with built-in example data first?
#    True  = ignore your files and use example data (good for a first try)
#    False = use YOUR files from STEP 2
USE_DEMO_DATA = True


# 👉 STEP 1 - What do you want to do? (pick ONE and write it between the quotes)
#    "fuzzy_compare"   -> match rows of TABLE_A and TABLE_B even when the names are
#                         spelled differently, e.g. "Apex Cement Co." vs "APEX CEMENT FACTORY"
#    "compare_tables"  -> match rows by a key that is EXACTLY the same in both tables
#                         (invoice no., item code, employee ID ...)
#    "compare_columns" -> compare one column with another column
#    "compare_rows"    -> compare one row with another row
#    "pivot"           -> build a pivot table from TABLE_A
TASK = "fuzzy_compare"


# 👉 STEP 2 - Your files: CSV, TXT, TSV, XLSX, XLSM, XLS, JSON or Parquet.
#    Power BI: on a visual click "..." -> Export data, then put that file here.
#    Windows paths: use forward slashes    "C:/Users/Me/Desktop/sales.xlsx"
#                   or put r in front      r"C:\Users\Me\Desktop\sales.xlsx"
#    "sheet" = the Excel sheet (tab) name. None = the first sheet. Ignored for CSV.
#    To compare two sheets of the SAME Excel file, use the same "file" with a different "sheet".
#    TABLE_B is not used by "pivot".
TABLE_A = {"file": "data/erp_suppliers.xlsx",      "sheet": None}
TABLE_B = {"file": "data/supplier_statement.csv",  "sheet": None}


# 👉 STEP 3 - KEY column(s): the column that says WHICH row is which
#    (customer name, supplier name, invoice no., item code ...).
#    Write the names EXACTLY as they appear in the header row of each file.
#    The names may be different in the two tables.
#    More than one column is allowed. Keep the same order on both sides:
#        KEY_A = ["Customer", "City"]
#        KEY_B = ["Client",   "Town"]
#    Used by: "fuzzy_compare" and "compare_tables"
KEY_A = ["Supplier Name"]
KEY_B = ["Vendor"]


# 👉 STEP 4 - Which columns should be checked once rows are matched?
#    Format:  {"column name in TABLE_A": "column name in TABLE_B", ...}
#    Leave it empty  {}  to check every column that has the same name in both tables.
#    Used by: "fuzzy_compare" and "compare_tables"
COMPARE_COLUMNS = {"Amount": "Total", "City": "City"}


# 👉 STEP 5 - Fuzzy settings (only for "fuzzy_compare")
SIMILARITY_THRESHOLD = 85     # 0-100. Pairs scoring below this are NOT matched.
                              # 85-90 is a good start. Lower = more matches, but more wrong ones.
REVIEW_BELOW = 95             # Matches scoring below this get "Needs review" (yellow in Excel).

MATCH_METHOD = "token_sort"   # "token_sort" -> ignores word order. Best for names.
                              #                 "Steel Blue River" = "Blue River Steel"
                              # "token_set"  -> OK when one side has extra words.
                              #                 "Apex Cement" vs "Apex Cement Factory Baghdad"
                              # "simple"     -> letter-by-letter. Best for codes/IDs with typos.
                              # "partial"    -> one text is inside the other.

IGNORE_WORDS = ["co", "company", "ltd", "llc", "inc", "the"]
                              # Words removed before comparing (upper/lower case doesn't matter).
                              # [] = remove nothing.

MUST_MATCH_EXACTLY = {}       # Optional safety rule: only pair rows that ALSO have exactly
                              # the same value in these columns.
                              # Example: {"Invoice Date": "Date"}   or   {"City": "City"}
                              # {} = off


# 👉 STEP 6 - General options
NUMBER_TOLERANCE = 0.0        # Numbers count as "same" if they differ by no more than this.
                              # Example: 0.01 to ignore rounding differences.
IGNORE_CASE = True            # True = "ERBIL" and "erbil" count as the same text.


# 👉 STEP 7 - Pivot settings (only for "pivot", uses TABLE_A)
PIVOT_ROWS = ["City"]         # fields down the left side
PIVOT_COLUMNS = ["Category"]  # fields across the top          ([] = none)
PIVOT_VALUES = ["Amount"]     # fields to calculate            ([] = just count rows)
PIVOT_CALC = "sum"            # "sum", "mean", "count", "nunique", "min", "max" or "median"
PIVOT_TOTALS = True           # add Total row/column


# 👉 STEP 8 - Column comparison (only for "compare_columns")
COLUMN_A = "Supplier Name"    # a column in TABLE_A
COLUMN_B = "Vendor"           # a column in TABLE_B
                              # (to compare two columns of ONE table, point TABLE_B to the same file)
COMPARE_BY_POSITION = False   # True = also compare row 1 with row 1, row 2 with row 2 ...
                              # Only useful if both tables are sorted the same way.


# 👉 STEP 9 - Row comparison (only for "compare_rows")
#    "find_by" = "row_number" (first data row is 0) or the name of a column to search in
#    "value"   = the row number, or the value to look for in that column
#    Columns with the same name are compared, plus the pairs you set in COMPARE_COLUMNS (STEP 4).
ROW_A = {"find_by": "Supplier Name", "value": "Apex Cement Factory"}
ROW_B = {"find_by": "Vendor",        "value": "Apex Cement Fctory"}


# 👉 STEP 10 - Where to save the report (the folder is created if needed)
OUTPUT_FILE = "results/result.xlsx"


# ╔══════════════════════════════════════════════════════════════════════════╗
# ║                ⛔  NO NEED TO EDIT ANYTHING BELOW THIS LINE               ║
# ╚══════════════════════════════════════════════════════════════════════════╝

import difflib
import io
import os
import re
import sys
from datetime import datetime

import numpy as np
import pandas as pd

try:
    from rapidfuzz import fuzz, process
except ImportError:  # only needed for fuzzy matching
    fuzz = process = None


FILE_TYPES = ["csv", "tsv", "txt", "xlsx", "xlsm", "xls", "json", "parquet"]

AGG_FUNCS = {"Sum": "sum", "Average": "mean", "Count": "count", "Distinct count": "nunique",
             "Min": "min", "Max": "max", "Median": "median"}
NUMERIC_AGGS = {"sum", "mean", "median"}

MATCH_METHODS = {
    "token_sort": ("token_sort_ratio", "Names - ignores word order"),
    "token_set": ("token_set_ratio", "Names with extra words on one side"),
    "simple": ("ratio", "Codes / IDs with typos"),
    "partial": ("partial_ratio", "One text contained in the other"),
}


class SettingsError(Exception):
    """A problem the user can fix in the SETTINGS section."""


# =============================================================================
# LOADING
# =============================================================================
def _read_text(data: bytes, sep):
    """Read CSV/TXT/TSV. Tries common encodings, including Arabic Windows (cp1256)."""
    for enc in ("utf-8-sig", "cp1256", "latin-1"):
        try:
            return pd.read_csv(io.BytesIO(data), sep=sep, engine="python", encoding=enc)
        except UnicodeDecodeError:
            continue
    raise ValueError("Could not detect the text encoding of this file.")


def clean_df(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    df.columns = [str(c).strip() for c in df.columns]
    df = df.dropna(how="all").dropna(axis=1, how="all")
    return df.reset_index(drop=True)


def read_file_bytes(filename: str, data: bytes) -> dict:
    """Return {sheet_name: DataFrame}. Non-Excel files use the key ""."""
    ext = os.path.splitext(filename)[1].lower().lstrip(".")
    if ext in ("csv", "txt"):
        return {"": clean_df(_read_text(data, sep=None))}   # auto-detects , ; | tab
    if ext == "tsv":
        return {"": clean_df(_read_text(data, sep="\t"))}
    if ext in ("xlsx", "xlsm", "xls"):
        sheets = pd.read_excel(io.BytesIO(data), sheet_name=None)
        return {str(s): clean_df(df) for s, df in sheets.items()}
    if ext == "json":
        return {"": clean_df(pd.read_json(io.BytesIO(data)))}
    if ext == "parquet":
        return {"": clean_df(pd.read_parquet(io.BytesIO(data)))}
    raise ValueError(f"Unsupported file type '.{ext}'. Supported: {', '.join(FILE_TYPES)}")


def load_table(cfg: dict, setting: str) -> pd.DataFrame:
    """Load one table from a TABLE_A / TABLE_B setting."""
    path = str(cfg.get("file", "")).strip()
    if not path:
        raise SettingsError(f'{setting}["file"] is empty. Put the path to your file there.')
    if not os.path.isfile(path):
        raise SettingsError(f'File not found: "{path}"  (check {setting}["file"])\n'
                            f'  Current folder is: {os.getcwd()}')
    with open(path, "rb") as f:
        try:
            sheets = read_file_bytes(path, f.read())
        except ValueError as e:
            raise SettingsError(f"{setting}: {e}")
    sheet = cfg.get("sheet")
    if "" in sheets:
        return sheets[""]
    if sheet is None:
        return next(iter(sheets.values()))
    if str(sheet) not in sheets:
        raise SettingsError(f'Sheet "{sheet}" not found in "{path}" (check {setting}["sheet"]).\n'
                            f"  Sheets in this file: {list(sheets)}")
    return sheets[str(sheet)]


def demo_tables():
    """Example: suppliers in an ERP vs a supplier's statement, names typed differently."""
    a = pd.DataFrame({
        "Supplier Name": ["Northstar Trading Co.", "Blue River Steel Ltd", "Apex Cement Factory",
                          "Global Logistics LLC", "Silver Line Electronics", "Green Valley Foods",
                          "Summit Office Supplies", "Delta Fuel Services"],
        "City": ["Riverside", "Hillview", "Lakeside", "Riverside", "Hillview", "Lakeside",
                 "Riverside", "Hillview"],
        "Category": ["Trading", "Construction", "Construction", "Transport", "IT", "Food",
                     "Office", "Fuel"],
        "Amount": [12500, 48200.5, 31000, 9800, 15250, 7400, 2300, 18900],
    })
    b = pd.DataFrame({
        "Vendor": ["NORTHSTAR TRADING COMPANY", "Blue River Steel", "Apex Cement Fctory",
                   "Global Logistics", "Silverline Electronics", "Green Valley Food Co",
                   "Orion Medical Supplies", "Delta Fuels Services"],
        "City": ["Riverside", "Hillview", "Lakeside", "Riverside", "Hillview", "Hillview",
                 "Lakeside", "Hillview"],
        "Total": [12500, 48200.5, 31500, 9800, 15250, 7400, 5600, 18900],
    })
    return a, b


# =============================================================================
# SMALL HELPERS
# =============================================================================
def check_columns(df: pd.DataFrame, cols, table: str, setting: str):
    """Friendly error (with 'did you mean') if a column name is wrong."""
    available = [str(c) for c in df.columns]
    problems = []
    for c in cols:
        if c not in df.columns:
            guess = difflib.get_close_matches(str(c), available, n=1, cutoff=0.6)
            hint = f'  -> did you mean "{guess[0]}"?' if guess else ""
            problems.append(f'  - "{c}" is not a column in {table}{hint}')
    if problems:
        raise SettingsError(f"Problem in setting {setting}:\n" + "\n".join(problems)
                            + f"\n  Columns in {table}: {available}")


def norm_text(s: pd.Series, ignore_case: bool = False) -> pd.Series:
    """Comparable text: trims spaces, '1001.0' -> '1001', empty -> ''."""
    out = s.astype(str).str.strip().str.replace(r"\.0$", "", regex=True)
    out = out.where(s.notna(), "")
    return out.str.lower() if ignore_case else out


def values_equal(a, b, tol: float = 0.0, ignore_case: bool = False) -> np.ndarray:
    """Cell-by-cell equality. Numbers use a tolerance, text is trimmed."""
    a = pd.Series(a).reset_index(drop=True)
    b = pd.Series(b).reset_index(drop=True)
    both_empty = a.isna() & b.isna()
    one_empty = a.isna() ^ b.isna()
    num_a = pd.to_numeric(a, errors="coerce")
    num_b = pd.to_numeric(b, errors="coerce")
    both_num = num_a.notna() & num_b.notna()
    num_eq = (num_a - num_b).abs() <= tol
    txt_eq = norm_text(a, ignore_case) == norm_text(b, ignore_case)
    return (both_empty | (both_num & num_eq) | (~both_num & ~one_empty & txt_eq)).to_numpy()


def flatten_columns(df: pd.DataFrame) -> pd.DataFrame:
    if isinstance(df.columns, pd.MultiIndex):
        df.columns = [" | ".join(str(x) for x in col if str(x) != "") for col in df.columns]
    else:
        df.columns = [str(c) for c in df.columns]
    return df


def file_row(i):
    """Row number as seen in Excel (header is row 1, first data row is 2)."""
    return np.asarray(i) + 2


# =============================================================================
# 1. PIVOT
# =============================================================================
def make_pivot(df, rows, columns, values, agg, totals=True, fill0=True) -> pd.DataFrame:
    data = df.copy()
    group_cols = rows + columns
    data[group_cols] = data[group_cols].astype(object).where(data[group_cols].notna(), "(blank)")

    if not values:                                   # nothing to calculate -> count rows
        data["Row count"] = 1
        values, agg = ["Row count"], "sum"
    elif agg in NUMERIC_AGGS:
        for v in values:
            data[v] = pd.to_numeric(data[v], errors="coerce")

    swap = not rows                                  # only columns chosen -> build then flip
    pt = pd.pivot_table(data, index=columns if swap else rows,
                        columns=None if swap else (columns or None),
                        values=values, aggfunc=agg, margins=totals, margins_name="Total",
                        fill_value=0 if fill0 else None)
    if swap:
        pt = pt.T
    if isinstance(pt.columns, pd.MultiIndex) and len(values) == 1:
        pt.columns = pt.columns.droplevel(0)
    pt = pt.reset_index()
    if swap:
        pt = pt.rename(columns={"index": "Value"})
    return flatten_columns(pt)


# =============================================================================
# 2. EXACT TABLE COMPARISON
# =============================================================================
def compare_schema(a: pd.DataFrame, b: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for c in dict.fromkeys(list(a.columns) + list(b.columns)):
        in_a, in_b = c in a.columns, c in b.columns
        rows.append({"Column": c,
                     "Status": "In both" if in_a and in_b else ("Only in A" if in_a else "Only in B"),
                     "Type in A": str(a[c].dtype) if in_a else "",
                     "Type in B": str(b[c].dtype) if in_b else ""})
    return pd.DataFrame(rows)


def compare_tables(a, b, keys, cols, tol=0.0, ignore_case=False) -> dict:
    """Match rows by key columns (same names in both tables) and report differences."""
    cols = [c for c in cols if c not in keys]
    a = a[keys + cols].copy()
    b = b[keys + cols].copy()
    for k in keys:                                   # 1001 / 1001.0 / " 1001" all match
        a[k] = norm_text(a[k], ignore_case)
        b[k] = norm_text(b[k], ignore_case)

    dup_a, dup_b = int(a.duplicated(keys).sum()), int(b.duplicated(keys).sum())
    m = a.merge(b, on=keys, how="outer", suffixes=(" (A)", " (B)"), indicator=True)

    only_a = m[m["_merge"] == "left_only"][keys + [f"{c} (A)" for c in cols]]
    only_a.columns = keys + cols
    only_b = m[m["_merge"] == "right_only"][keys + [f"{c} (B)" for c in cols]]
    only_b.columns = keys + cols
    both = m[m["_merge"] == "both"].reset_index(drop=True)

    parts, any_diff = [], np.zeros(len(both), dtype=bool)
    for c in cols:
        eq = values_equal(both[f"{c} (A)"], both[f"{c} (B)"], tol, ignore_case)
        any_diff |= ~eq
        d = both.loc[~eq, keys + [f"{c} (A)", f"{c} (B)"]].copy()
        if d.empty:
            continue
        d.columns = keys + ["Value A", "Value B"]
        d.insert(len(keys), "Column", c)
        parts.append(d)

    diffs = (pd.concat(parts, ignore_index=True) if parts
             else pd.DataFrame(columns=keys + ["Column", "Value A", "Value B"]))
    if not diffs.empty:
        diffs["Difference (B - A)"] = (pd.to_numeric(diffs["Value B"], errors="coerce")
                                       - pd.to_numeric(diffs["Value A"], errors="coerce"))
        diffs = diffs.sort_values(keys).reset_index(drop=True)

    return {"only_a": only_a.reset_index(drop=True), "only_b": only_b.reset_index(drop=True),
            "diffs": diffs, "matched": len(both), "matched_changed": int(any_diff.sum()),
            "dup_a": dup_a, "dup_b": dup_b}


def exact_compare(a, b, key_a, key_b, compare_cols, tol=0.0, ignore_case=False) -> dict:
    """Like compare_tables, but key/compare columns may have different names in A and B."""
    compare_cols = {ca: cb for ca, cb in compare_cols.items() if ca not in key_a}
    a2 = a[key_a + list(compare_cols)].copy()
    b2 = b[key_b + list(compare_cols.values())].copy()
    b2.columns = key_a + list(compare_cols)          # rename B's columns to A's names
    return compare_tables(a2, b2, key_a, list(compare_cols), tol, ignore_case)


# =============================================================================
# 3. FUZZY TABLE COMPARISON
# =============================================================================
def get_scorer(method: str):
    if fuzz is None:
        raise SettingsError("Fuzzy matching needs rapidfuzz. Install it with:  pip install rapidfuzz")
    if method not in MATCH_METHODS:
        raise SettingsError(f'MATCH_METHOD "{method}" is not valid. Use one of: {list(MATCH_METHODS)}')
    return getattr(fuzz, MATCH_METHODS[method][0])


def clean_for_matching(s: pd.Series, ignore_words=()) -> pd.Series:
    """lower case, remove punctuation and filler words, collapse spaces."""
    s = s.astype(object).where(s.notna(), "").astype(str).str.lower()
    s = s.str.replace(r"\.0$", "", regex=True)
    s = s.str.replace(r"[^\w\s]", " ", regex=True)   # \w keeps Arabic/Kurdish letters
    words = [w.strip().lower() for w in ignore_words if str(w).strip()]
    if words:
        s = s.str.replace(r"\b(" + "|".join(map(re.escape, words)) + r")\b", " ", regex=True)
    return s.str.replace(r"\s+", " ", regex=True).str.strip()


def _match_text(df, cols, ignore_words):
    out = clean_for_matching(df[cols[0]], ignore_words)
    for c in cols[1:]:
        out = out + " " + clean_for_matching(df[c], ignore_words)
    return out.str.strip().to_numpy(dtype=object)


def fuzzy_compare(a, b, key_a, key_b, compare_cols=None, threshold=85, method="token_sort",
                  must_match=None, ignore_words=(), review_below=95, tol=0.0,
                  ignore_case=True) -> dict:
    """
    Pair rows of A and B whose key text is similar (score 0-100).
    Each row is used at most once; the best-scoring pairs are chosen first.
    """
    scorer = get_scorer(method)
    must_match = must_match or {}
    a = a.reset_index(drop=True)
    b = b.reset_index(drop=True)
    if not compare_cols:
        skip = set(key_a) | set(key_b) | set(must_match) | set(must_match.values())
        compare_cols = {c: c for c in a.columns if c in b.columns and c not in skip}

    ta, tb = _match_text(a, key_a, ignore_words), _match_text(b, key_b, ignore_words)

    # Optional "must match exactly" groups: only rows in the same group are compared
    ga = np.full(len(a), "", dtype=object)
    gb = np.full(len(b), "", dtype=object)
    for ca, cb in must_match.items():
        ga = ga + "|" + norm_text(a[ca], ignore_case).to_numpy(dtype=object)
        gb = gb + "|" + norm_text(b[cb], ignore_case).to_numpy(dtype=object)

    best_a = np.zeros(len(a), dtype=int)
    best_a_idx = np.full(len(a), -1)
    best_b = np.zeros(len(b), dtype=int)
    best_b_idx = np.full(len(b), -1)
    cand_s, cand_i, cand_j = [], [], []

    for g in pd.unique(ga):
        ia = np.where((ga == g) & (ta != ""))[0]
        ib = np.where((gb == g) & (tb != ""))[0]
        if len(ia) == 0 or len(ib) == 0:
            continue
        m = process.cdist(ta[ia].tolist(), tb[ib].tolist(), scorer=scorer,
                          dtype=np.uint8, workers=-1)
        # second check with spaces removed: "Silver Line" = "Silverline"
        m2 = process.cdist([x.replace(" ", "") for x in ta[ia]], [x.replace(" ", "") for x in tb[ib]],
                           scorer=fuzz.ratio, dtype=np.uint8, workers=-1)
        m = np.maximum(m, m2)
        # closest candidate for every row (used to explain unmatched rows)
        ba = m.argmax(axis=1)
        best_a[ia], best_a_idx[ia] = m[np.arange(len(ia)), ba], ib[ba]
        bb = m.argmax(axis=0)
        best_b[ib], best_b_idx[ib] = m[bb, np.arange(len(ib))], ia[bb]
        r, c = np.nonzero(m >= threshold)
        cand_s.append(m[r, c].astype(int))
        cand_i.append(ia[r])
        cand_j.append(ib[c])

    # One-to-one assignment: highest scores first
    used_a = np.zeros(len(a), dtype=bool)
    used_b = np.zeros(len(b), dtype=bool)
    pairs = []
    if cand_s:
        s, i, j = np.concatenate(cand_s), np.concatenate(cand_i), np.concatenate(cand_j)
        for k in np.lexsort((j, i, -s)):
            if not used_a[i[k]] and not used_b[j[k]]:
                used_a[i[k]] = used_b[j[k]] = True
                pairs.append((int(i[k]), int(j[k]), int(s[k])))
    pairs.sort()
    pi = np.array([p[0] for p in pairs], dtype=int)
    pj = np.array([p[1] for p in pairs], dtype=int)
    ps = np.array([p[2] for p in pairs], dtype=int)

    # --- Matches table ---
    matches = pd.DataFrame({"Row in A file": file_row(pi), "Row in B file": file_row(pj)})
    for k in key_a:
        matches[f"{k} (A)"] = a[k].iloc[pi].to_numpy()
    for k in key_b:
        matches[f"{k} (B)"] = b[k].iloc[pj].to_numpy()
    raw_a = np.array([" ".join(x) for x in zip(*[norm_text(a[k], True).iloc[pi] for k in key_a])]) if len(pi) else np.array([])
    raw_b = np.array([" ".join(x) for x in zip(*[norm_text(b[k], True).iloc[pj] for k in key_b])]) if len(pj) else np.array([])
    matches["Similarity %"] = ps
    matches["Match type"] = np.where(raw_a == raw_b, "Identical",
                                     np.where(ps == 100, "Same after cleaning", "Similar"))
    matches["Needs review"] = np.where(ps < review_below, "Yes", "")
    key_cols = [c for c in matches.columns if c.endswith(" (A)") or c.endswith(" (B)")]

    # --- Value differences for matched pairs ---
    parts, n_diff = [], np.zeros(len(pairs), dtype=int)
    for ca, cb in compare_cols.items():
        va = a[ca].iloc[pi].reset_index(drop=True)
        vb = b[cb].iloc[pj].reset_index(drop=True)
        eq = values_equal(va, vb, tol, ignore_case)
        n_diff += ~eq
        if (~eq).any():
            d = matches.loc[~eq, key_cols + ["Similarity %"]].copy()
            d["Column"] = ca if ca == cb else f"{ca} <-> {cb}"
            d["Value A"] = va[~eq].to_numpy()
            d["Value B"] = vb[~eq].to_numpy()
            parts.append(d)
    if compare_cols:
        matches["Values"] = np.where(n_diff == 0, "All same",
                                     np.char.add(n_diff.astype(str), " different"))
    diffs = (pd.concat(parts, ignore_index=True) if parts
             else pd.DataFrame(columns=key_cols + ["Similarity %", "Column", "Value A", "Value B"]))
    if not diffs.empty:
        diffs["Difference (B - A)"] = (pd.to_numeric(diffs["Value B"], errors="coerce")
                                       - pd.to_numeric(diffs["Value A"], errors="coerce"))

    # --- Unmatched rows, with the closest candidate so the user can judge ---
    def unmatched(df, used, best, best_idx, other, other_keys, label):
        idx = np.where(~used)[0]
        out = df.iloc[idx].copy()
        out.insert(0, f"Row in {label} file", file_row(idx))
        out[f"Closest in {'B' if label == 'A' else 'A'}"] = [
            " / ".join(str(other[k].iloc[best_idx[x]]) for k in other_keys) if best_idx[x] >= 0 else ""
            for x in idx]
        out["Closest %"] = [int(best[x]) if best_idx[x] >= 0 else None for x in idx]
        return out.reset_index(drop=True)

    only_a = unmatched(a, used_a, best_a, best_a_idx, b, key_b, "A")
    only_b = unmatched(b, used_b, best_b, best_b_idx, a, key_a, "B")

    return {"matches": matches, "diffs": diffs, "only_a": only_a, "only_b": only_b,
            "stats": {"Rows in A": len(a), "Rows in B": len(b), "Matched pairs": len(pairs),
                      "Matches needing review": int((ps < review_below).sum()),
                      "Matched with different values": int((n_diff > 0).sum()),
                      "Only in A (no match)": len(only_a), "Only in B (no match)": len(only_b)}}


# =============================================================================
# 4. COLUMNS AND ROWS
# =============================================================================
def column_profile(s: pd.Series) -> dict:
    num = pd.to_numeric(s, errors="coerce")
    is_numeric = s.notna().sum() > 0 and num.notna().sum() >= 0.8 * s.notna().sum()
    p = {"Rows": len(s), "Empty": int(s.isna().sum()), "Distinct values": int(s.nunique())}
    if is_numeric:
        p.update({"Sum": round(num.sum(), 4), "Average": round(num.mean(), 4),
                  "Min": num.min(), "Max": num.max()})
    else:
        mode = s.mode()
        p["Most common"] = mode.iloc[0] if not mode.empty else ""
    return p


def value_overlap(sa: pd.Series, sb: pd.Series, ignore_case=False):
    A = set(norm_text(sa.dropna(), ignore_case))
    B = set(norm_text(sb.dropna(), ignore_case))
    return sorted(A & B), sorted(A - B), sorted(B - A)


def rowwise_compare(sa, sb, name_a, name_b, tol=0.0, ignore_case=False) -> pd.DataFrame:
    n = min(len(sa), len(sb))
    sa, sb = sa.iloc[:n].reset_index(drop=True), sb.iloc[:n].reset_index(drop=True)
    out = pd.DataFrame({"Row": range(n), name_a: sa, name_b: sb})
    out["Same?"] = values_equal(sa, sb, tol, ignore_case)
    diff = pd.to_numeric(sb, errors="coerce") - pd.to_numeric(sa, errors="coerce")
    if diff.notna().any():
        out["Difference (B - A)"] = diff
    return out


def find_row(df: pd.DataFrame, by: str, value):
    """Return (row, message). by = 'row_number' or a column name."""
    if str(by).lower().replace(" ", "_") == "row_number":
        try:
            i = int(value)
        except (TypeError, ValueError):
            return None, "Row number must be a whole number."
        if not 0 <= i < len(df):
            return None, f"Row number must be between 0 and {len(df) - 1}."
        return df.iloc[i], ""
    target = str(value).strip().lower().removesuffix(".0")
    hits = df[norm_text(df[by], True) == target]
    if hits.empty:
        return None, f"No row where {by} = '{value}'."
    return hits.iloc[0], (f"{len(hits)} rows matched; using the first." if len(hits) > 1 else "")


def compare_two_rows(ra: pd.Series, rb: pd.Series, ignore_case=True, tol=0.0,
                     pairs: dict | None = None) -> pd.DataFrame:
    """Compare same-named columns, plus any extra {col_in_A: col_in_B} pairs."""
    mapping = {c: c for c in ra.index if c in rb.index}
    for ca, cb in (pairs or {}).items():
        if ca in ra.index and cb in rb.index:
            mapping[ca] = cb
    out = pd.DataFrame({"Column": [ca if ca == cb else f"{ca} <-> {cb}" for ca, cb in mapping.items()],
                        "Row A": [ra[ca] for ca in mapping],
                        "Row B": [rb[cb] for cb in mapping.values()]})
    out["Same?"] = values_equal(out["Row A"], out["Row B"], tol, ignore_case)
    return out


# =============================================================================
# 5. EXCEL OUTPUT
# =============================================================================
def write_excel(sheets: dict, target) -> None:
    """Write {sheet_name: DataFrame} to a path or BytesIO, with simple formatting."""
    from openpyxl.styles import Font, PatternFill
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="305496")
    review_fill = PatternFill("solid", fgColor="FFF2CC")

    with pd.ExcelWriter(target, engine="openpyxl") as writer:
        for name, df in sheets.items():
            sn = re.sub(r"[\[\]:*?/\\]", "_", str(name))[:31] or "Sheet"
            df.to_excel(writer, sheet_name=sn, index=False)
            ws = writer.sheets[sn]
            ws.freeze_panes = "A2"
            for cell in ws[1]:
                cell.font, cell.fill = header_font, header_fill
            for col in ws.iter_cols(max_row=min(ws.max_row, 1000)):
                longest = max((len(str(c.value)) for c in col if c.value is not None), default=8)
                ws.column_dimensions[col[0].column_letter].width = min(longest + 2, 60)
            if "Needs review" in df.columns:
                k = list(df.columns).index("Needs review")
                for row in ws.iter_rows(min_row=2):
                    if row[k].value == "Yes":
                        for c in row:
                            c.fill = review_fill


def to_excel_bytes(sheets: dict) -> bytes:
    buf = io.BytesIO()
    write_excel(sheets, buf)
    return buf.getvalue()


# =============================================================================
# 6. RUN FROM SETTINGS
# =============================================================================
def _need_list(value, setting):
    if not isinstance(value, (list, tuple)) or not value:
        raise SettingsError(f'{setting} must be a list with at least one column, e.g. ["Name"]')
    return list(value)


def run_from_settings():
    """Returns (sheets dict, summary lines)."""
    task = str(TASK).strip().lower()
    valid = ["fuzzy_compare", "compare_tables", "compare_columns", "compare_rows", "pivot"]
    if task not in valid:
        raise SettingsError(f'TASK "{TASK}" is not valid. Use one of: {valid}')

    if USE_DEMO_DATA:
        a, b = demo_tables()
        name_a, name_b = "Demo: ERP suppliers", "Demo: supplier statement"
    else:
        a = load_table(TABLE_A, "TABLE_A")
        b = load_table(TABLE_B, "TABLE_B") if task != "pivot" else None
        name_a, name_b = TABLE_A["file"], TABLE_B["file"]

    info = [("Task", task), ("Table A", name_a)]
    if task != "pivot":
        info.append(("Table B", name_b))
    info.append(("Run at", datetime.now().strftime("%Y-%m-%d %H:%M")))

    # ---------- PIVOT ----------
    if task == "pivot":
        check_columns(a, PIVOT_ROWS + PIVOT_COLUMNS + PIVOT_VALUES, "TABLE_A",
                      "PIVOT_ROWS / PIVOT_COLUMNS / PIVOT_VALUES")
        if not PIVOT_ROWS and not PIVOT_COLUMNS:
            raise SettingsError("Put at least one column in PIVOT_ROWS or PIVOT_COLUMNS.")
        if PIVOT_CALC not in AGG_FUNCS.values():
            raise SettingsError(f'PIVOT_CALC "{PIVOT_CALC}" is not valid. Use one of: {list(AGG_FUNCS.values())}')
        result = make_pivot(a, list(PIVOT_ROWS), list(PIVOT_COLUMNS), list(PIVOT_VALUES),
                            PIVOT_CALC, PIVOT_TOTALS)
        info += [("Rows", ", ".join(PIVOT_ROWS)), ("Columns", ", ".join(PIVOT_COLUMNS)),
                 ("Values", ", ".join(PIVOT_VALUES) or "row count"), ("Calculation", PIVOT_CALC),
                 ("Result size", f"{len(result)} rows")]
        return {"Summary": pd.DataFrame(info, columns=["Item", "Value"]), "Pivot": result}, info

    # ---------- TABLE COMPARISONS ----------
    if task in ("fuzzy_compare", "compare_tables"):
        key_a, key_b = _need_list(KEY_A, "KEY_A"), _need_list(KEY_B, "KEY_B")
        if len(key_a) != len(key_b):
            raise SettingsError("KEY_A and KEY_B must have the same number of columns.")
        check_columns(a, key_a, "TABLE_A", "KEY_A")
        check_columns(b, key_b, "TABLE_B", "KEY_B")
        cols = dict(COMPARE_COLUMNS or {})
        if not cols:
            cols = {c: c for c in a.columns if c in b.columns and c not in key_a + key_b}
        check_columns(a, list(cols), "TABLE_A", "COMPARE_COLUMNS (left side)")
        check_columns(b, list(cols.values()), "TABLE_B", "COMPARE_COLUMNS (right side)")
        info += [("Key A", ", ".join(key_a)), ("Key B", ", ".join(key_b)),
                 ("Compared columns", ", ".join(f"{x} <-> {y}" for x, y in cols.items()) or "none"),
                 ("Number tolerance", NUMBER_TOLERANCE)]

        if task == "fuzzy_compare":
            mm = dict(MUST_MATCH_EXACTLY or {})
            check_columns(a, list(mm), "TABLE_A", "MUST_MATCH_EXACTLY (left side)")
            check_columns(b, list(mm.values()), "TABLE_B", "MUST_MATCH_EXACTLY (right side)")
            r = fuzzy_compare(a, b, key_a, key_b, cols, SIMILARITY_THRESHOLD, MATCH_METHOD, mm,
                              IGNORE_WORDS, REVIEW_BELOW, NUMBER_TOLERANCE, IGNORE_CASE)
            info += [("Method", MATCH_METHOD), ("Threshold", SIMILARITY_THRESHOLD),
                     ("Review below", REVIEW_BELOW)]
            info += list(r["stats"].items())
            return {"Summary": pd.DataFrame(info, columns=["Item", "Value"]),
                    "Matches": r["matches"], "Different values": r["diffs"],
                    "Only in A": r["only_a"], "Only in B": r["only_b"]}, info

        r = exact_compare(a, b, key_a, key_b, cols, NUMBER_TOLERANCE, IGNORE_CASE)
        info += [("Rows in A", len(a)), ("Rows in B", len(b)), ("Matched", r["matched"]),
                 ("Matched with different values", r["matched_changed"]),
                 ("Only in A", len(r["only_a"])), ("Only in B", len(r["only_b"]))]
        if r["dup_a"] or r["dup_b"]:
            info.append(("WARNING", f"Key not unique (duplicates A={r['dup_a']}, B={r['dup_b']}). "
                                    "Add another key column."))
        return {"Summary": pd.DataFrame(info, columns=["Item", "Value"]),
                "Different values": r["diffs"], "Only in A": r["only_a"],
                "Only in B": r["only_b"]}, info

    # ---------- COLUMNS ----------
    if task == "compare_columns":
        check_columns(a, [COLUMN_A], "TABLE_A", "COLUMN_A")
        check_columns(b, [COLUMN_B], "TABLE_B", "COLUMN_B")
        sa, sb = a[COLUMN_A], b[COLUMN_B]
        la, lb = f"A: {COLUMN_A}", f"B: {COLUMN_B}"
        prof = pd.DataFrame({la: pd.Series(column_profile(sa)), lb: pd.Series(column_profile(sb))})
        prof = prof.reset_index().rename(columns={"index": "Metric"})
        both, only_a, only_b = value_overlap(sa, sb, IGNORE_CASE)
        info += [("Column A", COLUMN_A), ("Column B", COLUMN_B), ("Values in both", len(both)),
                 ("Values only in A", len(only_a)), ("Values only in B", len(only_b))]
        sheets = {"Summary": pd.DataFrame(info, columns=["Item", "Value"]), "Column stats": prof,
                  "Values in both": pd.DataFrame({"Value": both}),
                  "Values only in A": pd.DataFrame({"Value": only_a}),
                  "Values only in B": pd.DataFrame({"Value": only_b})}
        if COMPARE_BY_POSITION:
            rw = rowwise_compare(sa, sb, la, lb, NUMBER_TOLERANCE, IGNORE_CASE)
            info.append(("Rows that differ (by position)", int((~rw["Same?"]).sum())))
            sheets["Summary"] = pd.DataFrame(info, columns=["Item", "Value"])
            sheets["Row by row"] = rw
        return sheets, info

    # ---------- ROWS ----------
    for cfg, df, setting in ((ROW_A, a, "ROW_A"), (ROW_B, b, "ROW_B")):
        by = cfg.get("find_by", "row_number")
        if str(by).lower().replace(" ", "_") != "row_number":
            check_columns(df, [by], "TABLE_A" if setting == "ROW_A" else "TABLE_B",
                          f'{setting}["find_by"]')
    ra, msg_a = find_row(a, ROW_A.get("find_by", "row_number"), ROW_A.get("value"))
    rb, msg_b = find_row(b, ROW_B.get("find_by", "row_number"), ROW_B.get("value"))
    if ra is None:
        raise SettingsError(f"ROW_A: {msg_a}")
    if rb is None:
        raise SettingsError(f"ROW_B: {msg_b}")
    out = compare_two_rows(ra, rb, IGNORE_CASE, NUMBER_TOLERANCE, dict(COMPARE_COLUMNS or {}))
    info += [("Row A", f"{ROW_A.get('find_by')} = {ROW_A.get('value')}"),
             ("Row B", f"{ROW_B.get('find_by')} = {ROW_B.get('value')}"),
             ("Shared columns", len(out)), ("Columns that differ", int((~out["Same?"]).sum()))]
    for m in (msg_a, msg_b):
        if m:
            info.append(("Note", m))
    return {"Summary": pd.DataFrame(info, columns=["Item", "Value"]), "Row comparison": out}, info


def main():
    try:
        sys.stdout.reconfigure(errors="replace")      # avoid crashes on old Windows consoles
    except Exception:
        pass
    try:
        sheets, info = run_from_settings()
        folder = os.path.dirname(OUTPUT_FILE)
        if folder:
            os.makedirs(folder, exist_ok=True)
        write_excel(sheets, OUTPUT_FILE)
    except SettingsError as e:
        print("\n❌ Please fix the SETTINGS:\n" + str(e) + "\n")
        sys.exit(1)
    except PermissionError:
        print(f'\n❌ Could not save "{OUTPUT_FILE}". Is it open in Excel? Close it and run again.\n')
        sys.exit(1)

    print("\n✅ Done\n" + "-" * 50)
    for k, v in info:
        print(f"{k:<32} {v}")
    print("-" * 50 + f"\nReport saved to: {os.path.abspath(OUTPUT_FILE)}\n")


if __name__ == "__main__":
    main()