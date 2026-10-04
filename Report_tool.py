"""
report_tool.py - ADD-ON for compare_tool.py
===========================================

Adds four things WITHOUT changing compare_tool.py or data_tool.py:

  1. Date grouping      - totals per day / week / month / quarter / year (gaps filled, change vs previous)
  2. Excel charts       - real, editable Excel charts + a Summary sheet with the key numbers
  3. Run without editing - command line options, saved "profiles" (JSON) and a run log,
                           so Windows Task Scheduler can produce the same report every month
  4. Formula checks     - e.g. "Total must equal Unit Price x Quantity".
                           Mismatches are LISTED - your data is never overwritten.

It only USES compare_tool.py (loading, cleaning, reading numbers/dates). Keep both files in
the same folder. Delete this file and everything else keeps working exactly as before.

HOW TO USE
  A) Edit the SETTINGS below (👉 comments) and run:     python report_tool.py
  B) Save those settings as a reusable profile:         python report_tool.py --save-profile monthly_sales
  C) Run a profile (no editing needed):                 python report_tool.py --profile monthly_sales
     ...on a new file:                                  python report_tool.py --profile monthly_sales --a "D:/exports/october.xlsx"
  D) Run any compare_tool task from the command line:   python report_tool.py --task check_data --a sales.xlsx
                                                        python report_tool.py --task fuzzy_compare --a erp.xlsx --b statement.csv
  E) See every setting name you can use in a profile:   python report_tool.py --list-settings

SCHEDULE IT (Windows)
  1. Edit run_report.bat so it points to your profile.
  2. Run once in Command Prompt (example: 08:00 on the 1st of every month):
       schtasks /create /tn "Monthly report" /tr "C:\\path\\to\\run_report.bat" /sc monthly /d 1 /st 08:00
  Every run is written to logs/run_log.csv (OK or FAILED, with the reason).
"""

# ╔══════════════════════════════════════════════════════════════════════════╗
# ║                  ✏️  SETTINGS - EDIT ONLY THIS PART                       ║
# ╚══════════════════════════════════════════════════════════════════════════╝

# 👉 STEP 0 - Test with built-in example data first?  (True = ignore your files)
USE_DEMO_DATA = True


# 👉 STEP 1 - What to run
#    "report"  -> the automatic report from this add-on (STEPS 3-7)
#    Or any compare_tool.py task: "check_data", "clean_data", "fuzzy_compare",
#    "compare_tables", "compare_columns", "compare_rows", "pivot"
#    (those use the settings in compare_tool.py, or the "compare_tool" part of a profile)
TASK = "report"


# 👉 STEP 2 - Your file(s). Same format as in compare_tool.py.
#    Windows paths: "C:/Users/Me/Desktop/sales.xlsx"  or  r"C:\Users\Me\Desktop\sales.xlsx"
#    "sheet": Excel sheet name, None = first sheet.  TABLE_B is only used by compare tasks.
TABLE_A = {"file": "data/sales.xlsx", "sheet": None}
TABLE_B = {"file": "", "sheet": None}


# 👉 STEP 3 - Group by date (for "report")
DATE_COLUMN = "Invoice Date"  # the date column to group by. None = no time grouping
PERIOD = "month"              # "day", "week", "month", "quarter" or "year"
                              # This adds a column named Day / Week / Month / Quarter / Year
                              # to the data. In the pivot (STEP 5) just write "Period" to use it.


# 👉 STEP 4 - What to measure (for "report")
VALUE_COLUMNS = ["Amount"]    # number columns to calculate. [] = just count rows
CALC = "sum"                  # "sum", "mean", "count", "nunique", "min", "max" or "median"
CATEGORY_COLUMN = "Category"  # a column to break the numbers down by (branch, product, region ...)
                              # None = no breakdown
TOP_N = 8                     # show the biggest N categories; the rest are grouped as "Other"


# 👉 STEP 5 - Extra pivot table with chart (optional, for "report")
#    Same idea as the pivot in compare_tool.py. Write "Period" to use the date grouping from
#    STEP 3 - it follows PERIOD automatically (month, quarter ...).   PIVOT_ROWS = [] turns it off.
PIVOT_ROWS = ["Period"]
PIVOT_COLUMNS = ["City"]
PIVOT_VALUES = ["Amount"]     # [] = count rows
PIVOT_CALC = "sum"
PIVOT_TOTALS = True
PIVOT_CHART = "line"          # "column", "bar", "line" or "none"


# 👉 STEP 6 - Formula checks (optional, for "report")
#    Write column names in [square brackets]. Allowed: numbers and  +  -  *  /  ( )
#    "check"        = the value in your data
#    "equals"       = what it should be
#    "tolerance"    = allowed difference (0.01 ignores rounding)
#    "tolerance_pct"= OR allowed difference in percent (e.g. 1 = 1%)
#    Rows that do not match are LISTED in the report. Nothing is changed.
#    Examples:
#      {"name": "Total = Price x Qty", "check": "[Total]", "equals": "[Unit Price] * [Quantity]", "tolerance": 0.01}
#      {"name": "Net = Gross - Tax",   "check": "[Net]",   "equals": "[Gross] - [Tax]"}
#      {"name": "Tax is 15%",          "check": "[Tax]",   "equals": "[Net] * 0.15", "tolerance_pct": 1}
FORMULA_CHECKS = [
    {"name": "Amount = Unit Price x Quantity", "check": "[Amount]",
     "equals": "[Unit Price] * [Quantity]", "tolerance": 0.01},
]
SHOW_COLUMNS = ["Invoice No", "Supplier Name"]   # extra columns shown next to each mismatch


# 👉 STEP 7 - Clean the data first? (for "report")
#    True = use the cleaning settings from compare_tool.py (STEP 10 and 11 there),
#    or from the "compare_tool" part of a profile. Every change goes to the "Cleaning log" sheet.
USE_CLEANING = True
INCLUDE_DATA_SHEET = True     # also put the (cleaned) data in the report


# 👉 STEP 8 - Where to save
#    You can use these placeholders in the file name, so each run keeps its own file:
#      {date} = 2026-10-04    {month} = 2026-10    {time} = 081500    {profile} = profile name
OUTPUT_FILE = "results/{profile}_report_{date}.xlsx"
LOG_FOLDER = "logs"           # run_log.csv is written here
PROFILES_FOLDER = "profiles"  # where --save-profile puts profiles


# ╔══════════════════════════════════════════════════════════════════════════╗
# ║                ⛔  NO NEED TO EDIT ANYTHING BELOW THIS LINE               ║
# ╚══════════════════════════════════════════════════════════════════════════╝

import argparse
import csv
import json
import os
import re
import sys
import time
import traceback
from datetime import datetime

import numpy as np
import pandas as pd

try:
    import compare_tool as core          # the core engine - used, never modified
except ImportError:
    print("❌ report_tool.py needs compare_tool.py in the same folder.")
    sys.exit(1)

SettingsError = core.SettingsError

PERIODS = {"day": ("D", "Day"), "week": ("W-SUN", "Week"), "month": ("M", "Month"),
           "quarter": ("Q", "Quarter"), "year": ("Y", "Year")}
CALCS = {"sum", "mean", "count", "nunique", "min", "max", "median"}
CHART_TYPES = {"column", "bar", "line", "none"}
CORE_TASKS = ["check_data", "clean_data", "fuzzy_compare", "compare_tables", "compare_columns",
              "compare_rows", "pivot"]
# These always come from report_tool settings / command line when running a core task
CONTROLLED_BY_ADDON = {"TASK", "TABLE_A", "TABLE_B", "OUTPUT_FILE", "USE_DEMO_DATA"}
FORMULA_KEYS = {"name", "check", "equals", "tolerance", "tolerance_pct"}
MAX_CHART_SERIES = 12


def settings_names(path: str) -> list:
    """Names of the settings in a tool's SETTINGS section (everything above the 'no need to edit' line)."""
    with open(path, encoding="utf-8") as f:
        text = f.read().split("NO NEED TO EDIT")[0]
    return list(dict.fromkeys(re.findall(r"^([A-Z][A-Z0-9_]*)\s*=", text, flags=re.M)))


ADDON_SETTINGS = settings_names(__file__)
CORE_SETTINGS = settings_names(core.__file__)


def current_settings() -> dict:
    return {k: globals()[k] for k in ADDON_SETTINGS}


# =============================================================================
# DEMO DATA (a year of sales with typical problems)
# =============================================================================
def demo_sales() -> pd.DataFrame:
    rng = np.random.default_rng(11)
    n = 420
    suppliers = ["Northstar Trading", "Blue River Steel", "Apex Cement", "Global Logistics",
                 "Silver Line Electronics", "Green Valley Foods", "Summit Office Supplies",
                 "Delta Fuel Services", "Orion Medical", "Harbor Tools", "Peak Textiles", "Cedar Furniture"]
    cats = {"Construction": 900, "IT": 450, "Food": 35, "Fuel": 120, "Office": 25,
            "Transport": 300, "Furniture": 220, "Textiles": 60, "Medical": 150}
    names = list(cats)
    dates = pd.Timestamp("2025-10-01") + pd.to_timedelta(rng.integers(0, 365, n), unit="D")
    season = 1 + 0.35 * np.sin((dates.month.to_numpy() - 3) / 12 * 2 * np.pi)
    cat = rng.choice(names, n, p=[.2, .15, .15, .12, .1, .1, .08, .05, .05])
    price = np.round([cats[c] * rng.uniform(0.7, 1.3) for c in cat], 2)
    qty = np.maximum(1, np.round(rng.integers(1, 20, n) * season)).astype(int)
    amount = np.round(price * qty, 2)
    disc = rng.choice(n, 7, replace=False)
    amount[disc] = np.round(amount[disc] * 0.9, 2)                     # 7 rows with an unrecorded discount
    df = pd.DataFrame({
        "Invoice No": [f"INV-{i:04d}" for i in range(1, n + 1)],
        "Supplier Name": rng.choice(suppliers, n),
        "City": rng.choice(["Riverside", "Hillview", "Lakeside", "Brookfield"], n),
        "Category": cat, "Unit Price": price, "Quantity": qty,
        "Amount": amount.astype(object),
        "Invoice Date": pd.Series(list(dates), dtype=object),   # object, so pandas can't re-parse the text dates
    })
    for i in rng.choice(n, 30, replace=False):                         # numbers typed as text
        df.at[i, "Amount"] = f"{float(df.at[i, 'Amount']):,.2f}"
    for i in rng.choice(n, 40, replace=False):                         # dates typed as text
        df.at[i, "Invoice Date"] = df.at[i, "Invoice Date"].strftime("%d/%m/%Y")
    for i in rng.choice(n, 15, replace=False):                         # spelling variants
        df.at[i, "City"] = df.at[i, "City"].upper() + " "
    return df


# =============================================================================
# VALIDATION
# =============================================================================
def _check_cols(df, cols, setting):
    available = [str(c) for c in df.columns]
    bad = [c for c in cols if c not in df.columns]
    if bad:
        lines = [f'  - "{c}" is not a column in the data{core._suggest(c, available)}' for c in bad]
        raise SettingsError(f"Problem in setting {setting}:\n" + "\n".join(lines)
                            + f"\n  Columns available: {available}")


def validate_report_settings(s: dict, df: pd.DataFrame, period_col):
    if s["DATE_COLUMN"]:
        _check_cols(df, [s["DATE_COLUMN"]], "DATE_COLUMN")
        if s["PERIOD"] not in PERIODS:
            raise SettingsError(f'PERIOD "{s["PERIOD"]}" is not valid. Use one of: {list(PERIODS)}')
    for name in ("VALUE_COLUMNS", "PIVOT_ROWS", "PIVOT_COLUMNS", "PIVOT_VALUES", "SHOW_COLUMNS"):
        if not isinstance(s[name], (list, tuple)):
            raise SettingsError(f'{name} must be a list like ["Column"] ([] = none).')
    _check_cols(df, s["VALUE_COLUMNS"], "VALUE_COLUMNS")
    _check_cols(df, s["SHOW_COLUMNS"], "SHOW_COLUMNS")
    for name in ("CALC", "PIVOT_CALC"):
        if s[name] not in CALCS:
            raise SettingsError(f'{name} "{s[name]}" is not valid. Use one of: {sorted(CALCS)}')
    if s["CATEGORY_COLUMN"]:
        _check_cols(df, [s["CATEGORY_COLUMN"]], "CATEGORY_COLUMN")
    if not isinstance(s["TOP_N"], int) or s["TOP_N"] < 1:
        raise SettingsError("TOP_N must be a whole number, 1 or more.")
    pivot_cols = list(s["PIVOT_ROWS"]) + list(s["PIVOT_COLUMNS"]) + list(s["PIVOT_VALUES"])
    if pivot_cols:
        extra = "" if period_col else "  (the period column only exists when DATE_COLUMN is set)"
        try:
            _check_cols(df, pivot_cols, "PIVOT_ROWS / PIVOT_COLUMNS / PIVOT_VALUES")
        except SettingsError as e:
            raise SettingsError(str(e) + extra)
    if s["PIVOT_CHART"] not in CHART_TYPES:
        raise SettingsError(f'PIVOT_CHART "{s["PIVOT_CHART"]}" is not valid. Use one of: {sorted(CHART_TYPES)}')
    if not isinstance(s["FORMULA_CHECKS"], (list, tuple)):
        raise SettingsError("FORMULA_CHECKS must be a list of checks ([] = none).")
    for n, chk in enumerate(s["FORMULA_CHECKS"], 1):
        if not isinstance(chk, dict) or "check" not in chk or "equals" not in chk:
            raise SettingsError(f'FORMULA_CHECKS item {n} needs "check" and "equals", e.g. '
                                '{"check": "[Total]", "equals": "[Price] * [Qty]"}')
        unknown = set(chk) - FORMULA_KEYS
        if unknown:
            raise SettingsError(f"FORMULA_CHECKS item {n}: unknown key(s) {sorted(unknown)}. "
                                f"Use: {sorted(FORMULA_KEYS)}")


# =============================================================================
# 4. FORMULA CHECKS (flag only)
# =============================================================================
def eval_formula(df: pd.DataFrame, formula, label: str) -> pd.Series:
    """Safely calculate '[Unit Price] * [Quantity]' for every row."""
    if isinstance(formula, (int, float)):
        return pd.Series(float(formula), index=df.index)
    formula = str(formula)
    cols = re.findall(r"\[([^\]]+)\]", formula)
    rest = re.sub(r"\[[^\]]+\]", "", formula)
    if not re.fullmatch(r"[\d\s+\-*/().]*", rest):
        raise SettingsError(f'{label}: "{formula}" - only [Column names], numbers and + - * / ( ) are allowed.')
    _check_cols(df, cols, label)
    env, expr = {}, formula
    for i, c in enumerate(dict.fromkeys(cols)):
        env[f"c{i}"] = core.parse_numbers(df[c], core.DECIMAL_SEPARATOR)[0].to_numpy(dtype=float)
        expr = expr.replace(f"[{c}]", f"c{i}")
    if not re.fullmatch(r"[c\d\s+\-*/().]*", expr):
        raise SettingsError(f'{label}: "{formula}" could not be read.')
    try:
        with np.errstate(all="ignore"):
            result = eval(expr, {"__builtins__": {}}, env)     # only numbers/operators can reach here
    except (SyntaxError, ZeroDivisionError, TypeError) as e:
        raise SettingsError(f'{label}: "{formula}" is not a valid formula ({e}).')
    return pd.Series(np.broadcast_to(np.asarray(result, dtype=float), len(df)).copy(), index=df.index)


def run_formula_checks(df: pd.DataFrame, checks, show_cols):
    summary, problems = [], []
    for n, chk in enumerate(checks, 1):
        name = chk.get("name") or f"{chk['check']} = {chk['equals']}"
        left = eval_formula(df, chk["check"], f'FORMULA_CHECKS item {n} "check"')
        right = eval_formula(df, chk["equals"], f'FORMULA_CHECKS item {n} "equals"')
        ok_vals = np.isfinite(left) & np.isfinite(right)
        diff = left - right
        bad = ok_vals & (diff.abs() > float(chk.get("tolerance", 0.01)))
        if "tolerance_pct" in chk:
            bad &= diff.abs() > right.abs() * float(chk["tolerance_pct"]) / 100
        summary.append({"Check": name, "Rows checked": int(ok_vals.sum()), "Mismatches": int(bad.sum()),
                        "Could not check (empty/unreadable)": int((~ok_vals).sum()),
                        "Total difference": round(float(diff[bad].sum()), 2)})
        for i in df.index[bad.to_numpy()]:
            row = {"Check": name, "Row in file": int(core.file_row(i))}
            row.update({c: df.at[i, c] for c in show_cols})
            row.update({"Value": left[i], "Expected": round(right[i], 4), "Difference": round(diff[i], 4),
                        "Difference %": round(diff[i] / right[i] * 100, 2) if right[i] else None})
            problems.append(row)
    return pd.DataFrame(summary), pd.DataFrame(problems)


# =============================================================================
# 1. DATE GROUPING + 2. REPORT TABLES
# =============================================================================
def period_label(p, period: str) -> str:
    if period == "week":
        return f"Week of {p.start_time:%Y-%m-%d}"
    if period == "quarter":
        return f"{p.year}-Q{p.quarter}"
    return str(p)


def _agg(groups, col, calc):
    return groups.size() if col is None else groups[col].agg(calc)


def build_report(df: pd.DataFrame, s: dict, source: str, profile: str):
    """Returns (sheets, chart specs, summary lines)."""
    info = [("Report generated", datetime.now().strftime("%Y-%m-%d %H:%M")),
            ("Profile", profile or "-"), ("Source", source)]
    sheets, charts, extra = {}, [], {}

    # ---- clean (uses compare_tool's cleaning settings) ----
    if s["USE_CLEANING"]:
        cfg = core.cleaning_config_from_settings()
        try:
            core.validate_cleaning(cfg, {"A": df})
        except SettingsError as e:
            raise SettingsError(f"{e}\n  (Cleaning settings come from compare_tool.py STEP 10-11, or the "
                                '"compare_tool" part of your profile. Or set USE_CLEANING = False.)')
        rows_before = len(df)
        df, log, probs = core.clean_table(df, cfg, "A")
        log_df, prob_df = core.logs_to_frames(log, probs)
        info.append(("Cleaning", f"{int(log_df['Cells changed'].sum())} cells changed, "
                                 f"{rows_before - len(df)} duplicate rows removed, {len(prob_df)} data problems"))
        extra = {"Cleaning log": log_df.drop(columns=["Table"]), "Data problems": prob_df.drop(columns=["Table"])}
    else:
        info.append(("Cleaning", "off"))
    df = df.copy()
    info.append(("Rows in report", len(df)))

    # ---- period column ----
    period_col, p = None, None
    date_col, period = s["DATE_COLUMN"], s["PERIOD"]
    if date_col and date_col in df.columns and period in PERIODS:
        freq, word = PERIODS[period]
        period_col = word if word not in df.columns else f"{date_col} ({word.lower()})"
        dates = df[date_col] if pd.api.types.is_datetime64_any_dtype(df[date_col]) \
            else core.parse_dates(df[date_col], core.DAY_FIRST)[0]
        p = dates.dt.to_period(freq)
        df[period_col] = [period_label(x, period) if pd.notna(x) else "(no date)" for x in p]
    if period_col:                       # "Period" (or Month/Quarter/...) in pivot settings -> the period column
        words = {"period"} | {w.lower() for _, w in PERIODS.values()}
        for name in ("PIVOT_ROWS", "PIVOT_COLUMNS"):
            s[name] = [period_col if (str(c).lower() in words and c not in df.columns) else c for c in s[name]]
    validate_report_settings(s, df, period_col)

    values = list(s["VALUE_COLUMNS"])
    for v in values:
        if not pd.api.types.is_numeric_dtype(df[v]):
            df[v] = core.parse_numbers(df[v], core.DECIMAL_SEPARATOR)[0]
    calc = s["CALC"]
    metric = values[0] if values else None
    metric_name = f"{metric} ({calc})" if metric else "Rows"
    fill0 = calc in ("sum", "count", "nunique")

    for v in values:
        col = df[v].dropna()
        info += [(f"{v} - total", round(float(col.sum()), 2)), (f"{v} - average", round(float(col.mean()), 2)
                                                                 if len(col) else ""),
                 (f"{v} - smallest", col.min() if len(col) else ""), (f"{v} - largest", col.max() if len(col) else "")]

    # ---- trend over time ----
    full = None
    if p is not None:
        valid = p.notna()
        no_date = int((~valid).sum())
        info += [("Date column", date_col), ("Grouped by", period)]
        if valid.any():
            d = dates[valid]
            info += [("From", d.min().date()), ("To", d.max().date())]
            full = pd.period_range(p[valid].min(), p[valid].max(), freq=PERIODS[period][0])
            g = df[valid.to_numpy()].groupby(p[valid])
            trend = pd.DataFrame(index=full)
            for v in values:
                trend[f"{v} ({calc})"] = _agg(g, v, calc).reindex(full)
                if fill0:
                    trend[f"{v} ({calc})"] = trend[f"{v} ({calc})"].fillna(0)
            trend["Rows"] = g.size().reindex(full).fillna(0).astype(int)
            base = trend[metric_name]
            trend["Change vs previous %"] = (base.pct_change(fill_method=None) * 100).replace([np.inf, -np.inf], np.nan).round(1)
            trend.insert(0, PERIODS[period][1], [period_label(x, period) for x in full])
            trend = trend.reset_index(drop=True)
            sheets["Trend"] = trend
            series = [trend.columns.get_loc(f"{v} ({calc})") + 1 for v in values] or [trend.columns.get_loc("Rows") + 1]
            charts.append({"sheet": "Trend", "kind": "line" if len(trend) >= 3 else "column",
                           "title": f"{metric_name} per {period}", "rows": len(trend), "cats": 1, "series": series})
            last = trend.iloc[-1]
            info += [(f"Latest {period}", last.iloc[0]), (f"{metric_name} in latest {period}", last[metric_name])]
            if len(trend) > 1 and pd.notna(last["Change vs previous %"]):
                info.append((f"Change vs previous {period}", f"{last['Change vs previous %']}%"))
            if d.max().normalize() < full[-1].end_time.normalize():
                info.append(("Note", f"The latest {period} is not complete - data runs until {d.max().date()}"))
        if no_date:
            info.append(("Rows without a valid date (not in the trend)", no_date))

    # ---- breakdown by category ----
    cat_col = s["CATEGORY_COLUMN"]
    mapped = None
    if cat_col:
        cat = df[cat_col].astype(object).where(df[cat_col].notna(), "(blank)")
        rank = _agg(df.groupby(cat), metric, calc).sort_values(ascending=False)
        top = list(rank.index[:s["TOP_N"]])
        mapped = cat.where(cat.isin(top), "Other")
        g = df.groupby(mapped)
        table = pd.DataFrame({f"{v} ({calc})": _agg(g, v, calc) for v in values})
        table["Rows"] = g.size()
        if metric and calc in ("sum", "count"):
            table[f"Share % (of {metric})"] = (table[metric_name] / table[metric_name].sum() * 100).round(1)
        else:
            table["Share % (of rows)"] = (table["Rows"] / table["Rows"].sum() * 100).round(1)
        order = top + (["Other"] if "Other" in table.index and "Other" not in top else [])
        table = table.reindex(order).reset_index().rename(columns={"index": cat_col})
        table.columns = [cat_col] + list(table.columns[1:])
        sheet = f"By {cat_col}"[:31]
        sheets[sheet] = table
        charts.append({"sheet": sheet, "kind": "bar", "title": f"{metric_name} by {cat_col}",
                       "rows": len(table), "cats": 1, "series": [2]})
        info += [(f"{cat_col} - distinct values", int(cat.nunique())),
                 (f"Largest {cat_col}", f"{table.iloc[0, 0]} ({table.iloc[0, -1]}%)")]

        if full is not None:
            valid = p.notna().to_numpy()
            tmp = pd.DataFrame({"_p": p[valid], "_c": mapped[valid],
                                "_v": df.loc[valid, metric] if metric else 1})
            pt = tmp.pivot_table(index="_p", columns="_c", values="_v",
                                 aggfunc=calc if metric else "sum").reindex(full)
            pt = pt[[c for c in order if c in pt.columns]]
            if fill0 or not metric:
                pt = pt.fillna(0)
            pt.insert(0, PERIODS[period][1], [period_label(x, period) for x in full])
            pt = pt.reset_index(drop=True)
            pt.columns = [str(c) for c in pt.columns]
            sheet = f"Trend by {cat_col}"[:31]
            sheets[sheet] = pt
            charts.append({"sheet": sheet, "kind": "stacked" if fill0 else "line",
                           "title": f"{metric_name} per {period} by {cat_col}", "rows": len(pt), "cats": 1,
                           "series": list(range(2, min(len(pt.columns), MAX_CHART_SERIES + 1) + 1))})

    # ---- extra pivot ----
    rows, cols = list(s["PIVOT_ROWS"]), list(s["PIVOT_COLUMNS"])
    if rows or cols:
        pv = core.make_pivot(df, rows, cols, list(s["PIVOT_VALUES"]), s["PIVOT_CALC"], s["PIVOT_TOTALS"])
        sheets["Pivot"] = pv
        if s["PIVOT_CHART"] != "none" and len(rows) == 1:
            n_rows = len(pv) - (1 if s["PIVOT_TOTALS"] else 0)
            last_col = len(pv.columns) - (1 if s["PIVOT_TOTALS"] and cols else 0)
            series = list(range(2, min(last_col, MAX_CHART_SERIES + 1) + 1))
            if n_rows > 0 and series:
                charts.append({"sheet": "Pivot", "kind": s["PIVOT_CHART"], "title": "Pivot",
                               "rows": n_rows, "cats": 1, "series": series})

    # ---- formula checks ----
    if s["FORMULA_CHECKS"]:
        fsum, fprob = run_formula_checks(df, s["FORMULA_CHECKS"], list(s["SHOW_COLUMNS"]))
        sheets["Formula checks"] = fsum
        sheets["Formula mismatches"] = fprob
        for _, r in fsum.iterrows():
            info.append((f"Check: {r['Check']}", f"{r['Mismatches']} mismatch(es) in {r['Rows checked']} rows "
                                                 f"({r['Could not check (empty/unreadable)']} could not be checked)"))

    sheets.update(extra)
    if s["INCLUDE_DATA_SHEET"]:
        sheets["Data"] = df
    info = [(k, round(v, 2) if isinstance(v, float) else v) for k, v in info]
    summary = pd.DataFrame(info, columns=["Item", "Value"])
    return {"Summary": summary, **sheets}, charts, info


# =============================================================================
# 2. NATIVE EXCEL CHARTS (added after compare_tool.write_excel)
# =============================================================================
def add_charts_and_formats(path: str, charts: list, skip_format=("Data", "Cleaning log", "Data problems")):
    from openpyxl import load_workbook
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.utils import get_column_letter

    wb = load_workbook(path)
    for ws in wb.worksheets:                       # readable numbers, one format per column
        if ws.title in skip_format:
            continue
        for col in ws.iter_cols(min_row=2):
            nums = [c for c in col if isinstance(c.value, (int, float)) and not isinstance(c.value, bool)]
            if not nums:
                continue
            if ws.title == "Summary":                 # mixed column: format each cell on its own
                for c in nums:
                    if isinstance(c.value, float) and not float(c.value).is_integer():
                        c.number_format = "#,##0.00"
                    elif abs(c.value) >= 1000:
                        c.number_format = "#,##0"
                continue
            has_decimals = any(isinstance(c.value, float) and not float(c.value).is_integer() for c in nums)
            if has_decimals:
                fmt = "#,##0.00"
            elif any(abs(c.value) >= 1000 for c in nums):
                fmt = "#,##0"
            else:
                continue
            for c in nums:
                c.number_format = fmt
    for ws in wb.worksheets:                       # print nicely: landscape, fit to page width
        if ws.title not in skip_format:
            ws.page_setup.orientation = "landscape"
            ws.sheet_properties.pageSetUpPr.fitToPage = True
            ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    if "Summary" in wb.sheetnames:
        wb["Summary"].column_dimensions["A"].width = 46
        wb["Summary"].column_dimensions["B"].width = 40

    for spec in charts:
        ws = wb[spec["sheet"][:31]]
        kind = spec["kind"]
        if kind == "line":
            ch = LineChart()
        else:
            ch = BarChart()
            ch.type = "bar" if kind == "bar" else "col"
            if kind == "stacked":
                ch.grouping, ch.overlap = "stacked", 100
        for c in spec["series"]:
            ch.add_data(Reference(ws, min_col=c, min_row=1, max_row=spec["rows"] + 1), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=spec["cats"], min_row=2, max_row=spec["rows"] + 1))
        ch.title = spec["title"]
        ch.height, ch.width = 9, 20
        ch.x_axis.delete = ch.y_axis.delete = False          # show axis labels (hidden by default in openpyxl 3.1+)
        ch.y_axis.number_format = "#,##0"
        ch.y_axis.majorGridlines = ch.y_axis.majorGridlines   # keep gridlines
        if kind == "line":
            for ser in ch.series:
                ser.smooth = False                           # straight lines: don't invent curves between points
        if len(spec["series"]) == 1:
            ch.legend = None
        if kind == "bar":
            ch.x_axis.scaling.orientation = "maxMin"        # biggest at the top
        ws.add_chart(ch, f"{get_column_letter(ws.max_column + 2)}2")
    wb.save(path)


# =============================================================================
# 3. PROFILES, COMMAND LINE, RUN LOG
# =============================================================================
def profile_path(name: str) -> str:
    if os.path.isfile(name):
        return name
    base = name if name.endswith(".json") else f"{name}.json"
    return base if os.path.dirname(base) else os.path.join(PROFILES_FOLDER, base)


def save_profile(name: str) -> str:
    path = profile_path(name) if os.path.dirname(name) else os.path.join(
        PROFILES_FOLDER, name if name.endswith(".json") else f"{name}.json")
    data = {
        "_about": ["Profile for report_tool.py. Run:  python report_tool.py --profile " +
                   os.path.splitext(os.path.basename(path))[0],
                   "'report_tool' = settings from report_tool.py; 'compare_tool' = settings from compare_tool.py.",
                   "Same names and meaning as the 👉 comments in those files. Edit values here as needed.",
                   "TASK / TABLE_A / TABLE_B / OUTPUT_FILE / USE_DEMO_DATA are always taken from 'report_tool'."],
        "report_tool": current_settings(),
        "compare_tool": {k: getattr(core, k) for k in CORE_SETTINGS if k not in CONTROLLED_BY_ADDON},
    }
    folder = os.path.dirname(path)
    if folder:
        os.makedirs(folder, exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2, ensure_ascii=False, default=str)
    return path


def load_profile(name: str):
    path = profile_path(name)
    if not os.path.isfile(path):
        raise SettingsError(f'Profile not found: "{name}" (looked for "{path}").')
    try:
        with open(path, encoding="utf-8-sig") as f:
            data = json.load(f)
    except json.JSONDecodeError as e:
        raise SettingsError(f'Profile "{path}" is not valid JSON: {e.msg} (line {e.lineno}, column {e.colno}).\n'
                            "  Tip: text needs \"double quotes\", True/False are written true/false, None is null.")
    bad_sections = set(data) - {"_about", "report_tool", "compare_tool"}
    if bad_sections:
        raise SettingsError(f'Profile "{path}": unknown section(s) {sorted(bad_sections)}. '
                            'Use "report_tool" and "compare_tool".')
    for section, allowed in (("report_tool", ADDON_SETTINGS), ("compare_tool", CORE_SETTINGS)):
        for k in data.get(section, {}):
            if k not in allowed:
                raise SettingsError(f'Profile "{path}", section "{section}": "{k}" is not a setting'
                                    f"{core._suggest(k, allowed)}")
    core_part = {k: v for k, v in data.get("compare_tool", {}).items() if k not in CONTROLLED_BY_ADDON}
    return data.get("report_tool", {}), core_part, os.path.splitext(os.path.basename(path))[0]


def resolve_output(pattern: str, profile: str) -> str:
    now = datetime.now()
    try:
        return pattern.format(date=now.strftime("%Y-%m-%d"), month=now.strftime("%Y-%m"),
                              time=now.strftime("%H%M%S"), profile=profile or "manual")
    except (KeyError, IndexError, ValueError) as e:
        raise SettingsError(f'OUTPUT_FILE "{pattern}" has an unknown placeholder ({e}). '
                            "Use {date}, {month}, {time} or {profile}.")


def write_run_log(folder: str, row: dict):
    try:
        os.makedirs(folder, exist_ok=True)
        path = os.path.join(folder, "run_log.csv")
        new = not os.path.isfile(path)
        with open(path, "a", newline="", encoding="utf-8-sig" if new else "utf-8") as f:
            w = csv.DictWriter(f, fieldnames=list(row))
            if new:
                w.writeheader()
            w.writerow(row)
    except OSError as e:
        print(f"⚠ Could not write the run log ({e}). Is logs/run_log.csv open in Excel?")


def parse_args(argv):
    ap = argparse.ArgumentParser(description="Automatic reports + command-line runs for compare_tool.py")
    ap.add_argument("--profile", help="profile name or .json path (see --save-profile)")
    ap.add_argument("--task", help='"report" or a compare_tool task, e.g. check_data, fuzzy_compare')
    ap.add_argument("--a", help="file for TABLE_A"), ap.add_argument("--a-sheet", help="Excel sheet for TABLE_A")
    ap.add_argument("--b", help="file for TABLE_B"), ap.add_argument("--b-sheet", help="Excel sheet for TABLE_B")
    ap.add_argument("--out", help="output file (placeholders allowed)")
    ap.add_argument("--demo", action="store_true", help="use the built-in demo data")
    ap.add_argument("--save-profile", metavar="NAME", help="save the current settings as a profile and stop")
    ap.add_argument("--list-settings", action="store_true", help="show all setting names and stop")
    return ap.parse_args(argv)


def main(argv=None) -> int:
    try:
        sys.stdout.reconfigure(errors="replace")
    except Exception:
        pass
    args = parse_args(sys.argv[1:] if argv is None else argv)

    if args.list_settings:
        print("report_tool settings :", ", ".join(ADDON_SETTINGS))
        print("compare_tool settings:", ", ".join(k for k in CORE_SETTINGS if k not in CONTROLLED_BY_ADDON))
        return 0
    if args.save_profile:
        print(f"✅ Profile saved: {os.path.abspath(save_profile(args.save_profile))}")
        return 0

    t0 = time.time()
    s, profile, out, task = current_settings(), "", "", ""
    log_row = {"Started": datetime.now().strftime("%Y-%m-%d %H:%M:%S"), "Profile": "", "Task": "",
               "Input A": "", "Input B": "", "Output": "", "Status": "", "Seconds": "", "Message": ""}
    try:
        core_part = {}
        log_row["Profile"] = args.profile or "-"
        if args.profile:
            addon_part, core_part, profile = load_profile(args.profile)
            s.update(addon_part)
        if args.task:
            s["TASK"] = args.task
        if args.a:
            s["TABLE_A"], s["USE_DEMO_DATA"] = {"file": args.a, "sheet": args.a_sheet}, False
        if args.b:
            s["TABLE_B"] = {"file": args.b, "sheet": args.b_sheet}
        if args.out:
            s["OUTPUT_FILE"] = args.out
        if args.demo:
            s["USE_DEMO_DATA"] = True

        for k, v in core_part.items():           # profile values for compare_tool settings (in memory only)
            setattr(core, k, v)
        task = str(s["TASK"]).strip().lower()
        out = resolve_output(s["OUTPUT_FILE"], profile)
        log_row.update({"Profile": profile or "-", "Task": task, "Output": out,
                        "Input A": "demo" if s["USE_DEMO_DATA"] else s["TABLE_A"].get("file", ""),
                        "Input B": "" if task in ("report", "pivot") else
                        ("demo" if s["USE_DEMO_DATA"] else s["TABLE_B"].get("file", ""))})
        if os.path.dirname(out):
            os.makedirs(os.path.dirname(out), exist_ok=True)

        if task == "report":
            if s["USE_DEMO_DATA"]:
                df, source = demo_sales(), "Demo: sales"
            else:
                df, source = core.load_table(s["TABLE_A"], "TABLE_A"), s["TABLE_A"]["file"]
            sheets, charts, info = build_report(df, s, source, profile)
            core.write_excel(sheets, out)
            add_charts_and_formats(out, charts)
        elif task in CORE_TASKS:
            core.TASK, core.USE_DEMO_DATA = task, s["USE_DEMO_DATA"]
            core.TABLE_A, core.TABLE_B, core.OUTPUT_FILE = s["TABLE_A"], s["TABLE_B"], out
            sheets, info = core.run_from_settings()
            core.write_excel(sheets, out)
        else:
            raise SettingsError(f'TASK "{task}" is not valid. Use "report" or one of: {CORE_TASKS}')

        print("\n✅ Done\n" + "-" * 64)
        for k, v in info:
            print(f"{str(k)[:46]:<47} {v}")
        print("-" * 64 + f"\nReport saved to: {os.path.abspath(out)}\n")
        log_row.update({"Status": "OK", "Message": f"{len(sheets)} sheets"})
        code = 0
    except SettingsError as e:
        print("\n❌ Please fix the settings:\n" + str(e) + "\n")
        log_row.update({"Status": "FAILED", "Message": " ".join(str(e).split())[:300]})
        code = 1
    except PermissionError:
        msg = f'Could not save "{out}". Is it open in Excel? Close it and run again.'
        print("\n❌ " + msg + "\n")
        log_row.update({"Status": "FAILED", "Message": msg})
        code = 1
    except Exception as e:
        os.makedirs(s.get("LOG_FOLDER") or "logs", exist_ok=True)
        err_file = os.path.join(s.get("LOG_FOLDER") or "logs", "last_error.txt")
        with open(err_file, "w", encoding="utf-8") as f:
            f.write(traceback.format_exc())
        print(f"\n❌ Unexpected error: {e}\n   Details saved to {os.path.abspath(err_file)}\n")
        log_row.update({"Status": "FAILED", "Message": f"Unexpected error: {e}"[:300]})
        code = 2
    log_row["Seconds"] = round(time.time() - t0, 1)
    write_run_log(s.get("LOG_FOLDER") or "logs", log_row)
    return code


if __name__ == "__main__":
    sys.exit(main())